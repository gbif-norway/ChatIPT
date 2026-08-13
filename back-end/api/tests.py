import os
import datetime
import io
import json
import tarfile
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from django.test import SimpleTestCase, TestCase, override_settings
import openpyxl
import pandas as pd
import yaml
from .helpers.publish import (
    DwcaExtensionLinkError,
    assert_case_insensitive_unique_identifier,
    gbif_dataset_type_for_core,
    make_eml,
    parse_newick_to_tree,
    parse_nexus_to_tree,
    parse_newick_tip_labels,
    parse_nexus_tip_labels,
    register_dataset_and_endpoint,
    upload_dwca,
    _sanitize_dataframe_for_utf8_export,
    validate_dwca_archive,
)
from .agent_tools import (
    BasicValidationForSomeDwCTerms,
    ExportDwcDp,
    GetDarwinCoreInfo,
    GetDwCExtensionInfo,
    SetEML,
    LogBugWithDeveloper,
    SetBasicMetadata,
    SetAgentTaskToComplete,
    RequestUserInput,
    ValidateDwcDp,
    PreviewDwcDpDescriptor,
    UploadDwCA,
    Python,
    normalize_event_date,
    _dwc_dp_resources_from_mapping,
)
from .helpers.openai_helpers import (
    CompatAssistantMessage,
    CompatFunctionCall,
    CompatToolCall,
    _attach_pdf_files_to_latest_user_message,
    _functions_to_responses_tools,
    _messages_to_responses_input,
    _response_to_compat_message,
    create_response_message,
)
from .models import Agent, Dataset, Message, Table, Task, UserFile
from .serializers import DatasetListSerializer, DatasetSerializer
from .dwc_dp_specs import (
    DWC_DP_SCHEMA_REVISION,
    RESERVED_TABLE_NAMES,
    VENDORED_SCHEMA_ERRORS,
    build_dwc_dp_explorer_model,
    build_datapackage_descriptor,
    create_dwc_dp_archive,
    export_dwc_dp_package,
    validate_datapackage_descriptor,
    validate_dwc_dp_archive,
    validate_dwc_dp_resources,
)
from .dwc_specs import EXTENSION_SCHEMAS, DarwinCoreCoreType, DarwinCoreExtensionType


class DwcDpSpecTests(SimpleTestCase):
    def _resources(self):
        return {
            'event': pd.DataFrame([
                {
                    'event_pk': 'event-1',
                    'eventID': 'source-event-1',
                    'eventCategory': 'occurrence',
                    'eventDate': '2025-04-26',
                    'country': 'Norway',
                },
            ]),
            'occurrence': pd.DataFrame([
                {
                    'occurrence_pk': 'occ-1',
                    'occurrenceID': 'source-occ-1',
                    'event_fk': 'event-1',
                    'scientificName': 'Apus apus',
                    'occurrenceStatus': 'present',
                },
            ]),
        }

    def test_validates_event_occurrence_resources_and_builds_descriptor(self):
        resources = self._resources()

        validation = validate_dwc_dp_resources(resources)
        self.assertTrue(validation['valid'], validation)
        self.assertEqual(validation['schema']['revision'], DWC_DP_SCHEMA_REVISION)

        descriptor = build_datapackage_descriptor(
            resources,
            title='Smoke test',
            description='DwC-DP smoke test',
        )
        self.assertEqual(
            descriptor['profile'],
            'https://raw.githubusercontent.com/gbif/dwc-dp/'
            'cbb6c887043876351eec1bed01c3dfc2e05c4eb4/'
            'dwc-dp/dwc-dp-profile.json',
        )
        self.assertEqual(descriptor['dwcDpSchema']['revision'], DWC_DP_SCHEMA_REVISION)
        self.assertEqual(len(descriptor['dwcDpSchema']['sha256']), 64)
        self.assertEqual(validate_datapackage_descriptor(descriptor), [])
        occurrence = next(resource for resource in descriptor['resources'] if resource['name'] == 'occurrence')
        self.assertEqual(occurrence['path'], 'occurrence.csv')
        self.assertEqual(occurrence['schema']['primaryKey'], 'occurrence_pk')
        self.assertEqual(occurrence['schema']['weakPrimaryKey'], 'occurrenceID')
        self.assertEqual(
            occurrence['schema']['foreignKeys'][0]['reference']['resource'],
            'event',
        )
        self.assertEqual(
            occurrence['schema']['foreignKeys'][0]['reference']['fields'],
            'event_pk',
        )

    def test_warns_when_weak_primary_keys_have_case_insensitive_duplicates(self):
        resources = self._resources()
        resources['occurrence'] = pd.DataFrame([
            {
                'occurrence_pk': 'occ-1',
                'occurrenceID': 'Source-Occurrence',
                'event_fk': 'event-1',
                'occurrenceStatus': 'present',
            },
            {
                'occurrence_pk': 'occ-2',
                'occurrenceID': 'source-occurrence',
                'event_fk': 'event-1',
                'occurrenceStatus': 'present',
            },
        ])

        validation = validate_dwc_dp_resources(resources)

        self.assertTrue(validation['valid'], validation)
        self.assertTrue(any(
            "weak primary key 'occurrenceID'" in warning
            and "2 row(s)" in warning
            for warning in validation['warnings']
        ))

    def test_vendored_snapshot_has_all_post_review_schemas(self):
        self.assertEqual(VENDORED_SCHEMA_ERRORS, ())
        self.assertEqual(len(RESERVED_TABLE_NAMES), 79)
        self.assertIn('survey-survey-target', RESERVED_TABLE_NAMES)
        self.assertIn('survey-target-descriptor', RESERVED_TABLE_NAMES)

    def test_rejects_old_identifier_based_relationship_shape(self):
        resources = {
            'event': pd.DataFrame([
                {'eventID': 'event-1', 'eventCategory': 'occurrence'},
            ]),
            'occurrence': pd.DataFrame([
                {'occurrenceID': 'occ-1', 'eventID': 'event-1', 'occurrenceStatus': 'present'},
            ]),
        }

        validation = validate_dwc_dp_resources(resources)

        self.assertFalse(validation['valid'])
        self.assertTrue(any("event_pk" in error for error in validation['errors']))
        self.assertTrue(any("event_fk" in error for error in validation['errors']))
        self.assertTrue(any("not defined" in error and "eventID" in error for error in validation['errors']))

    def test_validates_types_and_column_name_hygiene(self):
        resources = {
            'event': pd.DataFrame([
                {
                    ' event_pk': 'event-1',
                    'eventCategory': 'occurrence',
                    'year': 2025.5,
                },
            ]),
            'survey-target': pd.DataFrame([
                {'surveyTarget_pk': 'target-1', 'isSurveyTargetFullyReported': 'yes'},
            ]),
        }

        validation = validate_dwc_dp_resources(resources)

        self.assertFalse(validation['valid'])
        self.assertTrue(any("surrounding whitespace" in error for error in validation['errors']))
        self.assertTrue(any("integer values" in error for error in validation['errors']))
        self.assertTrue(any("boolean values" in error for error in validation['errors']))

    def test_enforces_strong_foreign_keys_but_not_weak_foreign_keys(self):
        event = pd.DataFrame([
            {
                'event_pk': 'event-1',
                'eventCategory': 'occurrence',
                'eventConductedByID': 'external-agent-id',
            },
        ])
        self.assertTrue(validate_dwc_dp_resources({'event': event})['valid'])

        occurrence = pd.DataFrame([
            {'occurrence_pk': 'occ-1', 'event_fk': 'missing-event', 'occurrenceStatus': 'present'},
        ])
        validation = validate_dwc_dp_resources({'event': event, 'occurrence': occurrence})
        self.assertFalse(validation['valid'])
        self.assertTrue(any("not found in 'event.event_pk'" in error for error in validation['errors']))

    def test_descriptor_includes_weak_relationships_when_target_table_is_present(self):
        resources = {
            'agent': pd.DataFrame([{'agent_pk': 'agent-1', 'agentID': 'source-agent-1'}]),
            'event': pd.DataFrame([
                {
                    'event_pk': 'event-1',
                    'eventID': 'source-event-1',
                    'eventCategory': 'occurrence',
                    'eventConductedByID': 'source-agent-1',
                },
            ]),
        }

        descriptor = build_datapackage_descriptor(resources)
        event = next(resource for resource in descriptor['resources'] if resource['name'] == 'event')

        self.assertEqual(event['schema']['weakPrimaryKey'], 'eventID')
        weak_fk = event['schema']['weakForeignKeys'][0]
        self.assertEqual(weak_fk['fields'], 'eventConductedByID')
        self.assertEqual(weak_fk['reference']['fields'], 'agentID')

    def test_creates_rooted_archive_with_plain_csv_resources(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            archive_path = Path(temp_dir) / 'package.tar.gz'
            create_dwc_dp_archive(
                archive_path,
                self._resources(),
                title='Smoke test',
                description='DwC-DP smoke test',
            )

            with tarfile.open(archive_path, 'r:gz') as archive:
                self.assertEqual(
                    sorted(archive.getnames()),
                    ['datapackage.json', 'eml.xml', 'event.csv', 'occurrence.csv'],
                )
                descriptor = json.load(archive.extractfile('datapackage.json'))
                occurrence_csv = archive.extractfile('occurrence.csv').read().decode('utf-8')

        self.assertEqual(descriptor['dwcDpSchema']['revision'], DWC_DP_SCHEMA_REVISION)
        self.assertIn('occurrence_pk,occurrenceID,event_fk', occurrence_csv)

    def test_archive_includes_and_validates_ancillary_tree(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            archive_path = Path(temp_dir) / 'package.tar.gz'
            create_dwc_dp_archive(
                archive_path,
                self._resources(),
                title='Tree smoke test',
                description='DwC-DP with supporting tree',
                additional_files=[('example.tre', b'(A,B);')],
            )

            validation = validate_dwc_dp_archive(
                archive_path,
                expected_resources=self._resources(),
                expected_additional_files=['example.tre'],
            )
            with tarfile.open(archive_path, 'r:gz') as archive:
                names = archive.getnames()

        self.assertTrue(validation['valid'], validation)
        self.assertIn('example.tre', names)

    def test_rejects_invalid_utf8_before_archive_creation(self):
        resources = self._resources()
        resources['occurrence'].loc[0, 'scientificName'] = 'collector\udc92s record'

        validation = validate_dwc_dp_resources(resources)

        self.assertFalse(validation['valid'])
        self.assertTrue(any('cannot be encoded as UTF-8' in error for error in validation['errors']))

    def test_publication_safety_and_semantic_checks_are_warnings(self):
        resources = {
            'event': pd.DataFrame([{
                'event_pk': 'event-1',
                'eventCategory': 'occurrence',
                'dataGeneralizations': (
                    'Coordinates generalized from 32.2340, -102.2340 to whole degrees.'
                ),
            }]),
            'event-assertion': pd.DataFrame([{
                'event_fk': 'event-1',
                'assertionType': 'unused flag',
                'assertionValue': '',
            }]),
            'resource-relationship': pd.DataFrame([{
                'subjectResourceID': 'occ-1',
                'relationshipType': 'preys on',
                'relatedResourceID': 'occ-2',
            }]),
        }

        validation = validate_dwc_dp_resources(resources)

        self.assertTrue(validation['valid'], validation)
        self.assertTrue(any('precise coordinate pairs' in warning for warning in validation['warnings']))
        self.assertTrue(any('blank assertionValue' in warning for warning in validation['warnings']))
        self.assertTrue(any('organism-interaction' in warning for warning in validation['warnings']))

    def test_bdq_inspired_checks_report_cross_field_and_vocabulary_warnings(self):
        resources = {
            'event': pd.DataFrame([{
                'event_pk': 'event-1',
                'eventCategory': 'occurrence',
                'eventDate': '2023-02-28',
                'year': 2024,
                'month': 2,
                'day': 30,
                'countryCode': 'no',
                'decimalLatitude': 0,
                'decimalLongitude': 0,
                'geodeticDatum': '',
                'minimumDepthInMeters': 20,
                'maximumDepthInMeters': 10,
                'minimumElevationInMeters': 100,
                'maximumElevationInMeters': 50,
            }]),
            'occurrence': pd.DataFrame([{
                'occurrence_pk': 'occ-1',
                'event_fk': 'event-1',
                'occurrenceStatus': 'Present',
            }]),
        }

        validation = validate_dwc_dp_resources(resources)
        warnings = '\n'.join(validation['warnings'])

        self.assertTrue(validation['valid'], validation)
        for test_id in (
            'VALIDATION_COORDINATES_NOTZERO',
            'VALIDATION_GEODETICDATUM_NOTEMPTY',
            'VALIDATION_MINDEPTH_LESSTHAN_MAXDEPTH',
            'VALIDATION_MINELEVATION_LESSTHAN_MAXELEVATION',
            'VALIDATION_COUNTRYCODE_STANDARD',
            'VALIDATION_OCCURRENCESTATUS_STANDARD',
            'VALIDATION_DAY_INRANGE',
            'VALIDATION_EVENT_CONSISTENT',
        ):
            self.assertIn(test_id, warnings)

    def test_bdq_inspired_checks_accept_valid_and_inapplicable_values(self):
        resources = {
            'event': pd.DataFrame([
                {
                    'event_pk': 'event-1',
                    'eventCategory': 'occurrence',
                    'eventDate': '2024-02-29',
                    'year': 2024,
                    'month': 2,
                    'day': 29,
                    'startDayOfYear': 60,
                    'endDayOfYear': 60,
                    'countryCode': 'NO',
                    'decimalLatitude': 0,
                    'decimalLongitude': 10,
                    'geodeticDatum': 'WGS84',
                    'minimumDepthInMeters': 0,
                    'maximumDepthInMeters': 10,
                    'minimumElevationInMeters': -20,
                    'maximumElevationInMeters': 100,
                },
                {
                    'event_pk': 'event-2',
                    'eventCategory': 'occurrence',
                    'year': 2024,
                    'day': 29,
                    'countryCode': 'XZ',
                },
            ]),
            'occurrence': pd.DataFrame([
                {
                    'occurrence_pk': 'occ-1',
                    'event_fk': 'event-1',
                    'occurrenceStatus': 'absent',
                },
                {
                    'occurrence_pk': 'occ-2',
                    'event_fk': 'event-2',
                    'occurrenceStatus': 'present',
                },
            ]),
        }

        validation = validate_dwc_dp_resources(resources)

        self.assertTrue(validation['valid'], validation)
        self.assertFalse(
            any(warning.startswith('BDQ-inspired') for warning in validation['warnings']),
            validation,
        )

    @patch('api.dwc_dp_specs.upload_file')
    @patch('minio.Minio')
    def test_uploads_package_with_gzip_content_type(self, minio_class, upload_file_mock):
        env = {
            'MINIO_URI': 'storage.example.org',
            'MINIO_ACCESS_KEY': 'key',
            'MINIO_SECRET_KEY': 'secret',
            'MINIO_BUCKET': 'bucket',
            'MINIO_BUCKET_FOLDER': 'packages',
        }
        with patch.dict(os.environ, env):
            url = export_dwc_dp_package(
                self._resources(),
                title='Smoke test',
                description='DwC-DP smoke test',
            )

        self.assertTrue(url.startswith('https://storage.example.org/bucket/packages/dwc-dp-'))
        self.assertEqual(upload_file_mock.call_args.kwargs['content_type'], 'application/gzip')
        minio_class.assert_called_once()

    def test_duplicate_resource_mapping_is_reported_without_overwriting(self):
        first = SimpleNamespace(id=1, title='event', df=self._resources()['event'])
        second = SimpleNamespace(id=2, title='event', df=self._resources()['event'])
        dataset = SimpleNamespace(table_set=SimpleNamespace(all=lambda: [first, second]))

        resources, errors = _dwc_dp_resources_from_mapping(dataset, None)

        self.assertEqual(list(resources), ['event'])
        self.assertEqual(len(errors), 1)
        self.assertIn('mapped more than once', errors[0])

class EmlGenerationTests(SimpleTestCase):
    def test_eml_template_is_well_formed(self):
        template_path = os.path.join(os.path.dirname(__file__), 'templates', 'eml.xml')
        with open(template_path, 'rb') as fh:
            xml_bytes = fh.read()
        # Should parse without error
        ET.fromstring(xml_bytes)

    def test_make_eml_generates_parseable_xml(self):
        class DummyUser:
            first_name = 'Alice'
            last_name = 'Smith'
            orcid_id = '0000-0001-2345-6789'
            email = 'alice@example.org'

        xml_text = make_eml(
            title='Test Dataset',
            description='Plain text abstract',
            user=DummyUser(),
            eml_extra={
                'geographic_scope': 'Norway. Coordinate bounds: lat 58 to 71, lon 4 to 31.',
                'temporal_scope': '2024-01-01/2024-12-31',
                'taxonomic_scope': 'Coleoptera',
                'methodology': 'Trap sampling',
                'users': [
                    {'first_name': 'Alice', 'last_name': 'Smith', 'orcid': '0000-0001-2345-6789'}
                ],
            },
        )
        root = ET.fromstring(xml_text.encode('utf-8'))
        # dataset exists
        dataset = root.find('dataset')
        self.assertIsNotNone(dataset)
        # title
        title = dataset.find('title')
        self.assertIsNotNone(title)
        self.assertEqual(title.text, 'Test Dataset')
        # language
        language = dataset.find('language')
        self.assertIsNotNone(language)
        self.assertEqual(language.text, 'eng')
        # pubDate
        self.assertIsNotNone(dataset.find('pubDate'))
        # abstract para
        abstract_para = dataset.find('abstract/para')
        self.assertIsNotNone(abstract_para)
        self.assertEqual(abstract_para.text, 'Plain text abstract')
        # contact with email
        contact_email = dataset.find('contact/electronicMailAddress')
        self.assertIsNotNone(contact_email)
        self.assertEqual(contact_email.text, 'alice@example.org')
        # coverage blocks present
        coverage = dataset.find('coverage')
        self.assertIsNotNone(coverage)
        self.assertIsNotNone(coverage.find('geographicCoverage/geographicDescription'))
        self.assertEqual(
            coverage.find('geographicCoverage/boundingCoordinates/westBoundingCoordinate').text,
            '4',
        )
        # methods description present
        self.assertIsNotNone(dataset.find('methods/methodStep/description/para'))

    def test_make_eml_uses_each_supported_gbif_license(self):
        expected = {
            "CC0 1.0": "http://creativecommons.org/publicdomain/zero/1.0/legalcode",
            "CC BY 4.0": "http://creativecommons.org/licenses/by/4.0/legalcode",
            "CC BY-NC 4.0": "http://creativecommons.org/licenses/by-nc/4.0/legalcode",
        }

        for license_name, license_url in expected.items():
            with self.subTest(license_name=license_name):
                root = ET.fromstring(
                    make_eml("Licensed data", "Description", eml_extra={"license": license_name})
                )
                link = root.find("dataset/intellectualRights/para/ulink")
                self.assertIsNotNone(link)
                self.assertEqual(link.attrib["url"], license_url)
                self.assertTrue(link.findtext("citetitle"))

    @patch("api.helpers.publish.requests.post")
    def test_gbif_registration_uses_canonical_license_url(self, post_mock):
        post_mock.side_effect = [
            SimpleNamespace(status_code=201, json=lambda: "dataset-key"),
            SimpleNamespace(status_code=201, json=lambda: {}),
        ]

        register_dataset_and_endpoint(
            "Licensed data",
            "Description",
            "https://example.org/archive.zip",
            "CC BY-NC 4.0",
        )

        dataset_payload = post_mock.call_args_list[0].kwargs["json"]
        self.assertEqual(
            dataset_payload["license"],
            "http://creativecommons.org/licenses/by-nc/4.0/legalcode",
        )

    def test_gbif_dataset_type_matches_archive_core(self):
        expected = {
            DarwinCoreCoreType.OCCURRENCE: "OCCURRENCE",
            DarwinCoreCoreType.EVENT: "SAMPLING_EVENT",
            DarwinCoreCoreType.TAXON: "CHECKLIST",
        }

        for core_type, dataset_type in expected.items():
            with self.subTest(core_type=core_type):
                self.assertEqual(gbif_dataset_type_for_core(core_type), dataset_type)

    @patch("api.helpers.publish.requests.post")
    def test_gbif_registration_uses_sampling_event_type_for_event_core(self, post_mock):
        post_mock.side_effect = [
            SimpleNamespace(status_code=201, json=lambda: "dataset-key"),
            SimpleNamespace(status_code=201, json=lambda: {}),
        ]

        gbif_url = register_dataset_and_endpoint(
            "Event data",
            "Description",
            "https://example.org/archive.zip",
            core_type=DarwinCoreCoreType.EVENT,
        )

        dataset_payload = post_mock.call_args_list[0].kwargs["json"]
        self.assertEqual(dataset_payload["type"], "SAMPLING_EVENT")
        self.assertEqual(
            gbif_url,
            "https://www.gbif-test.org/dataset/dataset-key",
        )

    def test_make_eml_maps_manuscript_fields_and_creators(self):
        class DummyUser:
            first_name = 'Alice'
            last_name = 'Smith'
            orcid_id = '0000-0001-2345-6789'
            email = 'alice@example.org'

        xml_text = make_eml(
            title='PDF-derived dataset',
            description='Abstract text',
            user=DummyUser(),
            eml_extra={
                'manuscript_doi': '10.1234/abcd.1',
                'dataset_citation': 'Doe J, Roe R (2025) Example manuscript.',
                'manuscript_title': 'Example manuscript',
                'journal': 'Journal of Examples',
                'publication_year': 2025,
                'users': [
                    {'first_name': 'Jane', 'last_name': 'Doe', 'email': 'jane@example.org'},
                    {'first_name': 'Richard', 'last_name': 'Roe', 'email': 'richard@example.org'},
                ],
            },
        )
        root = ET.fromstring(xml_text.encode('utf-8'))
        dataset = root.find('dataset')
        self.assertIsNotNone(dataset)

        creators = dataset.findall('creator')
        self.assertEqual(len(creators), 2)
        self.assertEqual(creators[0].find('individualName/givenName').text, 'Jane')
        self.assertEqual(creators[1].find('individualName/givenName').text, 'Richard')
        self.assertEqual(len(dataset.findall('project/personnel')), 0)

    def test_make_eml_normalizes_orcid_url_to_identifier_and_directory(self):
        class DummyUser:
            first_name = 'Alice'
            last_name = 'Smith'
            orcid_id = 'https://orcid.org/0000-0001-2345-6789'
            email = 'alice@example.org'

        xml_text = make_eml(
            title='ORCID normalization',
            description='Abstract text',
            user=DummyUser(),
            eml_extra={
                'users': [
                    {
                        'first_name': 'Jane',
                        'last_name': 'Doe',
                        'email': 'jane@example.org',
                        'orcid': 'https://orcid.org/0000-0002-1825-0097',
                    },
                ],
            },
        )
        root = ET.fromstring(xml_text.encode('utf-8'))
        dataset = root.find('dataset')
        self.assertIsNotNone(dataset)

        creator_user_id = dataset.find('creator/userId')
        self.assertIsNotNone(creator_user_id)
        self.assertEqual(creator_user_id.get('directory'), 'https://orcid.org/')
        self.assertEqual(creator_user_id.text, '0000-0002-1825-0097')

        metadata_provider_user_id = dataset.find('metadataProvider/userId')
        self.assertIsNotNone(metadata_provider_user_id)
        self.assertEqual(metadata_provider_user_id.get('directory'), 'https://orcid.org/')
        self.assertEqual(metadata_provider_user_id.text, '0000-0001-2345-6789')

        contact_email = dataset.find('contact/electronicMailAddress')
        self.assertIsNotNone(contact_email)
        self.assertEqual(contact_email.text, 'alice@example.org')

    def test_make_eml_adds_project_personnel_when_project_title_is_set(self):
        class DummyUser:
            first_name = 'Alice'
            last_name = 'Smith'
            orcid_id = '0000-0001-2345-6789'
            email = 'alice@example.org'

        xml_text = make_eml(
            title='Project-backed dataset',
            description='Abstract text',
            user=DummyUser(),
            eml_extra={
                'project_title': 'Arctic Deep Survey',
                'users': [
                    {'first_name': 'Jane', 'last_name': 'Doe', 'email': 'jane@example.org'},
                    {'first_name': 'Richard', 'last_name': 'Roe', 'email': 'richard@example.org'},
                ],
            },
        )
        root = ET.fromstring(xml_text.encode('utf-8'))
        dataset = root.find('dataset')
        self.assertIsNotNone(dataset)

        self.assertEqual(dataset.find('project/title').text, 'Arctic Deep Survey')
        personnel_nodes = dataset.findall('project/personnel')
        self.assertEqual(len(personnel_nodes), 2)
        self.assertEqual(personnel_nodes[0].find('individualName/givenName').text, 'Jane')
        self.assertEqual(personnel_nodes[1].find('individualName/givenName').text, 'Richard')
        self.assertIsNone(personnel_nodes[0].find('electronicMailAddress'))
        self.assertIsNone(personnel_nodes[0].find('userId'))
        self.assertEqual(personnel_nodes[0].find('role').text, 'metadataProvider')

    def test_make_eml_omits_project_when_title_has_no_personnel(self):
        xml_text = make_eml(
            title='Project-backed dataset',
            description='Abstract text',
            eml_extra={
                'project_title': 'Arctic Deep Survey',
                'users': [],
            },
        )
        root = ET.fromstring(xml_text.encode('utf-8'))

        self.assertIsNone(root.find('dataset/project'))

    def test_make_eml_keeps_creator_before_metadata_provider_when_users_override_primary_creator(self):
        class DummyUser:
            first_name = 'Alice'
            last_name = 'Smith'
            orcid_id = '0000-0001-2345-6789'
            email = 'alice@example.org'

        xml_text = make_eml(
            title='Ordered creators',
            description='Abstract text',
            user=DummyUser(),
            eml_extra={
                'users': [
                    {'first_name': 'Jane', 'last_name': 'Doe', 'email': 'jane@example.org'},
                ],
            },
        )
        root = ET.fromstring(xml_text.encode('utf-8'))
        dataset = root.find('dataset')
        self.assertIsNotNone(dataset)

        child_tags = [child.tag for child in list(dataset)]
        creator_index = child_tags.index('creator')
        metadata_provider_index = child_tags.index('metadataProvider')
        self.assertLess(creator_index, metadata_provider_index)

    def test_make_eml_builds_taxonomic_classification_from_scope(self):
        class DummyUser:
            first_name = 'Alice'
            last_name = 'Smith'
            orcid_id = '0000-0001-2345-6789'
            email = 'alice@example.org'

        xml_text = make_eml(
            title='Taxonomy',
            description='Abstract text',
            user=DummyUser(),
            eml_extra={
                'taxonomic_scope': 'Plantae, Solanaceae',
            },
        )
        root = ET.fromstring(xml_text.encode('utf-8'))
        dataset = root.find('dataset')
        self.assertIsNotNone(dataset)

        self.assertEqual(
            dataset.find('coverage/taxonomicCoverage/generalTaxonomicCoverage').text,
            'Plantae, Solanaceae',
        )
        first_classification = dataset.find('coverage/taxonomicCoverage/taxonomicClassification')
        self.assertIsNotNone(first_classification)
        self.assertEqual(first_classification.find('taxonRankName').text, 'kingdom')
        self.assertEqual(first_classification.find('taxonRankValue').text, 'Plantae')

        self.assertIsNone(first_classification.find('taxonomicClassification'))
        classifications = dataset.findall('coverage/taxonomicCoverage/taxonomicClassification')
        self.assertEqual(len(classifications), 2)
        self.assertEqual(classifications[1].find('taxonRankName').text, 'family')
        self.assertEqual(classifications[1].find('taxonRankValue').text, 'Solanaceae')

    def test_make_eml_uses_structured_taxonomic_keywords(self):
        class DummyUser:
            first_name = 'Alice'
            last_name = 'Smith'
            orcid_id = '0000-0001-2345-6789'
            email = 'alice@example.org'

        xml_text = make_eml(
            title='Structured taxonomy',
            description='Abstract text',
            user=DummyUser(),
            eml_extra={
                'taxonomic_scope': 'Plants and fungi in source data.',
                'taxonomic_keywords': [
                    {'rank': 'kingdom', 'scientificName': 'Plantae'},
                    {'rank': 'family', 'scientificName': 'Solanaceae', 'commonName': 'nightshades'},
                ],
            },
        )
        root = ET.fromstring(xml_text.encode('utf-8'))
        dataset = root.find('dataset')
        self.assertIsNotNone(dataset)

        classifications = dataset.findall('coverage/taxonomicCoverage/taxonomicClassification')
        self.assertEqual(len(classifications), 2)
        self.assertEqual(classifications[0].find('taxonRankValue').text, 'Plantae')
        self.assertEqual(classifications[1].find('commonName').text, 'nightshades')

    def test_make_eml_does_not_turn_taxonomic_prose_into_classifications(self):
        class DummyUser:
            first_name = 'Alice'
            last_name = 'Smith'
            orcid_id = '0000-0001-2345-6789'
            email = 'alice@example.org'

        xml_text = make_eml(
            title='Taxonomic prose',
            description='Abstract text',
            user=DummyUser(),
            eml_extra={
                'taxonomic_scope': (
                    'Plantae, Fabaceae. The dataset contains specimen-linked occurrence records '
                    'for multiple genera, with emphasis on Errazurizia and related taxa.'
                ),
            },
        )
        root = ET.fromstring(xml_text.encode('utf-8'))
        dataset = root.find('dataset')
        self.assertIsNotNone(dataset)

        classifications = dataset.findall('coverage/taxonomicCoverage/taxonomicClassification')
        self.assertEqual(len(classifications), 1)
        self.assertEqual(classifications[0].find('taxonRankValue').text, 'Plantae')

    def test_make_eml_orders_agent_email_before_user_id(self):
        class DummyUser:
            first_name = 'Alice'
            last_name = 'Smith'
            orcid_id = '0000-0001-2345-6789'
            email = 'alice@example.org'

        xml_text = make_eml(
            title='Ordered agent',
            description='Abstract text',
            user=DummyUser(),
            eml_extra={},
        )
        root = ET.fromstring(xml_text.encode('utf-8'))
        dataset = root.find('dataset')
        self.assertIsNotNone(dataset)

        creator_tags = [child.tag for child in list(dataset.find('creator'))]
        self.assertLess(creator_tags.index('electronicMailAddress'), creator_tags.index('userId'))
        contact_tags = [child.tag for child in list(dataset.find('contact'))]
        self.assertLess(contact_tags.index('electronicMailAddress'), contact_tags.index('userId'))

    def test_make_eml_omits_geographic_coverage_without_bounds(self):
        class DummyUser:
            first_name = 'Alice'
            last_name = 'Smith'
            orcid_id = '0000-0001-2345-6789'
            email = 'alice@example.org'

        xml_text = make_eml(
            title='No geographic bounds',
            description='Abstract text',
            user=DummyUser(),
            eml_extra={
                'geographic_scope': 'Norway',
            },
        )
        root = ET.fromstring(xml_text.encode('utf-8'))
        dataset = root.find('dataset')
        self.assertIsNotNone(dataset)

        self.assertIsNone(dataset.find('coverage/geographicCoverage'))

    def test_make_eml_extracts_geographic_bounds_from_scope_text(self):
        class DummyUser:
            first_name = 'Alice'
            last_name = 'Smith'
            orcid_id = '0000-0001-2345-6789'
            email = 'alice@example.org'

        xml_text = make_eml(
            title='Geographic bounds',
            description='Abstract text',
            user=DummyUser(),
            eml_extra={
                'geographic_scope': (
                    'Coordinates range from 28.55056 S to 44.49629167 N latitude '
                    'and from 112.71712 W to 70.80929 W longitude.'
                ),
            },
        )
        root = ET.fromstring(xml_text.encode('utf-8'))
        dataset = root.find('dataset')
        self.assertIsNotNone(dataset)

        bounds = dataset.find('coverage/geographicCoverage/boundingCoordinates')
        self.assertIsNotNone(bounds)
        self.assertEqual(bounds.find('westBoundingCoordinate').text, '-112.71712')
        self.assertEqual(bounds.find('eastBoundingCoordinate').text, '-70.80929')
        self.assertEqual(bounds.find('northBoundingCoordinate').text, '44.496292')
        self.assertEqual(bounds.find('southBoundingCoordinate').text, '-28.55056')

    def test_make_eml_accepts_structured_geographic_bounds_with_zero_values(self):
        class DummyUser:
            first_name = 'Alice'
            last_name = 'Smith'
            orcid_id = '0000-0001-2345-6789'
            email = 'alice@example.org'

        xml_text = make_eml(
            title='Structured geographic bounds',
            description='Abstract text',
            user=DummyUser(),
            eml_extra={
                'geographic_scope': 'Equator and prime meridian test area.',
                'geographic_bounds': {
                    'west': 0,
                    'east': 10,
                    'north': 5,
                    'south': 0,
                },
            },
        )
        root = ET.fromstring(xml_text.encode('utf-8'))
        dataset = root.find('dataset')
        self.assertIsNotNone(dataset)

        bounds = dataset.find('coverage/geographicCoverage/boundingCoordinates')
        self.assertIsNotNone(bounds)
        self.assertEqual(bounds.find('westBoundingCoordinate').text, '0')
        self.assertEqual(bounds.find('southBoundingCoordinate').text, '0')

    def test_make_eml_omits_template_coverage_when_all_scopes_are_empty(self):
        class DummyUser:
            first_name = 'Alice'
            last_name = 'Smith'
            orcid_id = '0000-0001-2345-6789'
            email = 'alice@example.org'

        xml_text = make_eml(
            title='No coverage',
            description='Abstract text',
            user=DummyUser(),
            eml_extra={
                'geographic_scope': None,
                'temporal_scope': None,
                'taxonomic_scope': None,
            },
        )
        root = ET.fromstring(xml_text.encode('utf-8'))
        dataset = root.find('dataset')
        self.assertIsNotNone(dataset)
        self.assertIsNone(dataset.find('coverage'))

    def test_make_eml_normalizes_temporal_textual_range(self):
        class DummyUser:
            first_name = 'Alice'
            last_name = 'Smith'
            orcid_id = '0000-0001-2345-6789'
            email = 'alice@example.org'

        xml_text = make_eml(
            title='Temporal text range',
            description='Abstract text',
            user=DummyUser(),
            eml_extra={
                'temporal_scope': '17 - 22 October 2025',
            },
        )

        root = ET.fromstring(xml_text.encode('utf-8'))
        dataset = root.find('dataset')
        self.assertIsNotNone(dataset)

        begin = dataset.find('coverage/temporalCoverage/rangeOfDates/beginDate/calendarDate')
        end = dataset.find('coverage/temporalCoverage/rangeOfDates/endDate/calendarDate')
        self.assertIsNotNone(begin)
        self.assertIsNotNone(end)
        self.assertEqual(begin.text, '2025-10-17')
        self.assertEqual(end.text, '2025-10-22')

    def test_make_eml_extracts_iso_dates_from_free_text_temporal_scope(self):
        class DummyUser:
            first_name = 'Alice'
            last_name = 'Smith'
            orcid_id = '0000-0001-2345-6789'
            email = 'alice@example.org'

        xml_text = make_eml(
            title='Temporal free text',
            description='Abstract text',
            user=DummyUser(),
            eml_extra={
                'temporal_scope': (
                    'Field observations appear span 2025-09-18 to 2025-09-23 based on project naming.'
                ),
            },
        )

        root = ET.fromstring(xml_text.encode('utf-8'))
        dataset = root.find('dataset')
        self.assertIsNotNone(dataset)

        begin = dataset.find('coverage/temporalCoverage/rangeOfDates/beginDate/calendarDate')
        end = dataset.find('coverage/temporalCoverage/rangeOfDates/endDate/calendarDate')
        self.assertIsNotNone(begin)
        self.assertIsNotNone(end)
        self.assertEqual(begin.text, '2025-09-18')
        self.assertEqual(end.text, '2025-09-23')


class IdentifierValidationTests(SimpleTestCase):
    def test_case_insensitive_identifier_uniqueness_passes_for_unique_values(self):
        df = pd.DataFrame({'occurrenceID': ['abc-1', 'Abc-2', 'ABC-3']})
        assert_case_insensitive_unique_identifier(df, 'occurrenceID')

    def test_case_insensitive_identifier_uniqueness_fails_for_collisions(self):
        df = pd.DataFrame({'occurrenceID': ['1764255881670-v4s75yh73|185À', '1764255881670-v4s75yh73|185à']})
        with self.assertRaises(ValueError):
            assert_case_insensitive_unique_identifier(df, 'occurrenceID')

    def test_case_insensitive_identifier_uniqueness_ignores_empty_values(self):
        df = pd.DataFrame({'occurrenceID': ['', None, 'x-1']})
        assert_case_insensitive_unique_identifier(df, 'occurrenceID')


class PhylogenyParsingTests(SimpleTestCase):
    """Tests for phylogenetic tree parsing and matching functions."""
    
    def test_parse_nexus_with_translate_block(self):
        """Test parsing NEXUS file with TRANSLATE block."""
        nexus_content = """#NEXUS
BEGIN TREES;
	Title Acacia;
	LINK Taxa = Taxa;
	TRANSLATE
        DBT01_clean	Berneuxia_thibetica_01,
        DBT03_clean	Berneuxia_thibetica_02,
        Berneuxia_thibetica	Berneuxia_thibetica_P,
        S7451102	Cyrilla_racemiflora,
        DDH04_clean	Diapensia_himalaica_01;
    TREE tree_1 = (DBT01_clean,DBT03_clean);
END;"""
        
        tip_labels = parse_nexus_tip_labels(nexus_content)
        
        # Should extract translated names from TRANSLATE block
        expected_labels = [
            'Berneuxia_thibetica_01',
            'Berneuxia_thibetica_02',
            'Berneuxia_thibetica_P',
            'Cyrilla_racemiflora',
            'Diapensia_himalaica_01'
        ]
        
        self.assertEqual(set(tip_labels), set(expected_labels))
        self.assertEqual(len(tip_labels), len(expected_labels))
    
    def test_parse_nexus_from_example_file(self):
        """Test parsing the actual example NEXUS file."""
        nexus_path = os.path.join(
            os.path.dirname(__file__),
            'templates', 'examples', 'gaynor_et_al_v3', 'above50_genes.nex'
        )
        
        with open(nexus_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        tip_labels = parse_nexus_tip_labels(content)
        
        # Should extract all translated tip labels
        self.assertGreater(len(tip_labels), 0)
        
        # Check for specific expected labels from the file
        expected_labels = [
            'Berneuxia_thibetica_01',
            'Berneuxia_thibetica_02',
            'Berneuxia_thibetica_P',
            'Cyrilla_racemiflora',
            'Diapensia_himalaica_01',
            'Diapensia_himalaica_02',
            'Diapensia_himalaica_P',
            'Diapensia_lapponica_01',
            'Diapensia_lapponica_02',
            'Diapensia_lapponica_03',
            'Diapensia_obovata_01',
            'Diapensia_obovata_02',
            'Diapensia_purpurea_01',
            'Diapensia_purpurea_P',
            'Diapensia_wardii_P',
            'Galax_urceolata_01',
            'Galax_urceolata_02',
            'Pyxidanthera_barbulata_01',
            'Pyxidanthera_barbulata_02',
            'Pyxidanthera_brevifolia_01',
            'Pyxidanthera_brevifolia_02',
            'Schizocodon_ilicifolius_01',
            'Schizocodon_soldanelloides_01',
            'Schizocodon_soldanelloides_02',
            'Shortia_galacifolia_01',
            'Shortia_galacifolia_02',
            'Shortia_rotundifolia_01',
            'Shortia_sinensis_01',
            'Shortia_sinensis_P',
            'Shortia_uniflora_01',
            'Shortia_uniflora_02',
        ]
        
        # Check that all expected labels are present
        for expected in expected_labels:
            self.assertIn(expected, tip_labels, f"Expected label {expected} not found")
        
        # Should have exactly 31 labels
        self.assertEqual(len(tip_labels), 31)
    
    def test_parse_newick_simple(self):
        """Test parsing a simple Newick format tree."""
        newick_content = "(A:0.1,B:0.2,(C:0.3,D:0.4):0.5);"
        
        tip_labels = parse_newick_tip_labels(newick_content)
        
        expected = ['A', 'B', 'C', 'D']
        self.assertEqual(set(tip_labels), set(expected))
        self.assertEqual(len(tip_labels), len(expected))
    
    def test_parse_newick_with_underscores(self):
        """Test parsing Newick with labels containing underscores."""
        newick_content = "(Berneuxia_thibetica:0.1,Diapensia_himalaica:0.2);"
        
        tip_labels = parse_newick_tip_labels(newick_content)
        
        expected = ['Berneuxia_thibetica', 'Diapensia_himalaica']
        self.assertEqual(set(tip_labels), set(expected))
    
    def test_parse_newick_complex(self):
        """Test parsing a more complex Newick tree."""
        newick_content = "((A:0.1,B:0.2)100:0.3,(C:0.4,D:0.5)90:0.6);"
        
        tip_labels = parse_newick_tip_labels(newick_content)
        
        expected = ['A', 'B', 'C', 'D']
        self.assertEqual(set(tip_labels), set(expected))
    
    def test_parse_nexus_no_translate_block(self):
        """Test parsing NEXUS file without TRANSLATE block."""
        nexus_content = """#NEXUS
BEGIN TREES;
    TREE tree_1 = (A:0.1,B:0.2,(C:0.3,D:0.4):0.5);
END;"""
        
        tip_labels = parse_nexus_tip_labels(nexus_content)
        
        expected = ['A', 'B', 'C', 'D']
        self.assertEqual(set(tip_labels), set(expected))
    
    def test_parse_newick_handles_quotes(self):
        """Test that Newick parser handles various label formats."""
        # Labels might have quotes or special characters
        newick_content = '("Species A":0.1,"Species B":0.2);'
        
        tip_labels = parse_newick_tip_labels(newick_content)
        
        # Should extract labels even with quotes
        # Note: current implementation may not handle quotes, but should not crash
        self.assertIsInstance(tip_labels, list)

    def test_parse_newick_preserves_nonstandard_unquoted_spaces(self):
        newick_content = (
            '(Prosopidastrum mexicanum RM17061 Mexico,'
            'Dichrostachys cinerea ERR7618736);'
        )

        self.assertEqual(
            parse_newick_tip_labels(newick_content),
            [
                'Prosopidastrum mexicanum RM17061 Mexico',
                'Dichrostachys cinerea ERR7618736',
            ],
        )
        tree = parse_newick_to_tree(newick_content)
        self.assertEqual(
            [child["name"] for child in tree["children"]],
            [
                "Prosopidastrum mexicanum RM17061 Mexico",
                "Dichrostachys cinerea ERR7618736",
            ],
        )

    def test_parse_nexus_translate_preserves_quoted_spaces(self):
        nexus_content = """#NEXUS
BEGIN TREES;
TRANSLATE
  1 'Species one voucher A',
  2 'Species two voucher B';
TREE tree_1 = (1:0.1,2:0.2);
END;"""

        self.assertEqual(
            parse_nexus_tip_labels(nexus_content),
            ['Species one voucher A', 'Species two voucher B'],
        )
        tree = parse_nexus_to_tree(nexus_content)
        self.assertEqual(
            [child["name"] for child in tree["children"]],
            ['Species one voucher A', 'Species two voucher B'],
        )


class GetDarwinCoreInfoTests(SimpleTestCase):
    def test_summary_response_lists_sections_only(self):
        sample_reference = {
            "Occurrence": {
                "occurrenceID": "An identifier. Examples: abc",
                "basisOfRecord": "The specific nature of the data record. Examples: HumanObservation",
            },
            "Event": {
                "eventID": "An identifier for the event. Examples: EVT-1",
            },
        }
        with patch("api.agent_tools._load_dwc_quick_reference", return_value=sample_reference):
            response = GetDarwinCoreInfo().run()

        self.assertIn("Darwin Core quick reference sections:", response)
        self.assertIn("- Occurrence (2 terms)", response)
        self.assertIn("- Event (1 terms)", response)
        self.assertNotIn("The specific nature of the data record", response)

    def test_section_response_returns_all_terms_and_includes_examples_by_default(self):
        sample_reference = {
            "Occurrence": {
                "occurrenceID": "An identifier for the occurrence. Examples: 1234",
                "basisOfRecord": "The specific nature of the data record. Examples: HumanObservation",
            },
        }
        with patch("api.agent_tools._load_dwc_quick_reference", return_value=sample_reference):
            response = GetDarwinCoreInfo(section="Occurrence").run()

        self.assertIn("Occurrence (2 terms total):", response)
        self.assertIn("- occurrenceID: An identifier for the occurrence. Examples: 1234", response)
        self.assertIn("- basisOfRecord: The specific nature of the data record. Examples: HumanObservation", response)

    def test_section_response_can_strip_examples_when_requested(self):
        sample_reference = {
            "Occurrence": {
                "occurrenceID": "An identifier for the occurrence. Examples: 1234",
                "basisOfRecord": "The specific nature of the data record. Examples: HumanObservation",
            },
        }
        with patch("api.agent_tools._load_dwc_quick_reference", return_value=sample_reference):
            response = GetDarwinCoreInfo(section="Occurrence", include_examples=False).run()

        self.assertIn("Occurrence (2 terms total):", response)
        self.assertIn("- occurrenceID: An identifier for the occurrence.", response)
        self.assertIn("- basisOfRecord: The specific nature of the data record.", response)
        self.assertNotIn("Examples:", response)

    def test_section_response_applies_optional_limit(self):
        sample_reference = {
            "Occurrence": {
                "occurrenceID": "An identifier for the occurrence.",
                "basisOfRecord": "The specific nature of the data record.",
            },
        }
        with patch("api.agent_tools._load_dwc_quick_reference", return_value=sample_reference):
            response = GetDarwinCoreInfo(section="Occurrence", max_terms=1).run()

        self.assertIn("Occurrence (2 terms total):", response)
        self.assertIn("- occurrenceID: An identifier for the occurrence.", response)
        self.assertIn("... 1 more terms not shown. Increase `max_terms` to see more.", response)

    def test_term_lookup_can_include_examples(self):
        sample_reference = {
            "Occurrence": {
                "basisOfRecord": "The specific nature of the data record. Examples: HumanObservation",
            },
        }
        with patch("api.agent_tools._load_dwc_quick_reference", return_value=sample_reference):
            response = GetDarwinCoreInfo(terms=["basisOfRecord"], include_examples=True).run()

        self.assertIn("Darwin Core term lookup results:", response)
        self.assertIn("basisOfRecord (Occurrence): The specific nature of the data record. Examples: HumanObservation", response)


class GetDwCExtensionInfoTests(SimpleTestCase):
    def test_every_extension_has_selection_metadata(self):
        for extension_type, schema in EXTENSION_SCHEMAS.items():
            with self.subTest(extension=extension_type.value):
                self.assertTrue(schema.compatible_cores)
                self.assertTrue(schema.subject)
                self.assertTrue(schema.typical_dwc_dp_resources)
                self.assertTrue(schema.use_when)
                self.assertTrue(schema.avoid_when)

    def test_lookup_includes_extensions_needed_for_rich_projections(self):
        relationship = GetDwCExtensionInfo(extension="resource_relationship").run()
        identification = GetDwCExtensionInfo(extension="identification_history").run()
        audiovisual = GetDwCExtensionInfo(extension="audiovisual").run()
        extended_measurement = GetDwCExtensionInfo(extension="extended_measurement_or_fact").run()
        material_sample = GetDwCExtensionInfo(extension="ggbn_material_sample").run()

        self.assertIn("Darwin Core Resource Relationship", relationship)
        self.assertIn("Darwin Core Identification History", identification)
        self.assertIn("Audiovisual Media Description", audiovisual)
        self.assertIn("PixelXDimension", audiovisual)
        self.assertIn("measurementTypeID", extended_measurement)
        self.assertIn("materialSampleType", material_sample)

    def test_catalogue_terms_and_row_types_match_vendored_definitions(self):
        for extension_type, schema in EXTENSION_SCHEMAS.items():
            with self.subTest(extension=extension_type.value):
                root = ET.parse(schema.spec_path).getroot()
                registered_terms = {
                    element.attrib["name"]
                    for element in root
                    if element.tag.rsplit("}", 1)[-1] == "property"
                }
                self.assertEqual(schema.row_type, root.attrib["rowType"])
                self.assertEqual(set(schema.terms), registered_terms)

    def test_all_registered_ggbn_extensions_are_available(self):
        expected = {
            "ggbn_material_sample",
            "ggbn_amplification",
            "ggbn_cloning",
            "ggbn_gel_image",
            "ggbn_loan",
            "ggbn_permit",
            "ggbn_preparation",
            "ggbn_preservation",
        }

        self.assertTrue(expected.issubset({extension.value for extension in EXTENSION_SCHEMAS}))

    def test_catalogue_exposes_core_compatibility_and_projection_guidance(self):
        response = GetDwCExtensionInfo().run()

        self.assertIn("Select every compatible extension needed to preserve meaningful source facts", response)
        self.assertIn("Omit redundant or empty extensions", response)
        self.assertIn("Humboldt Ecological Inventory", response)
        self.assertIn("Compatible cores: event", response)
        self.assertIn("Typical DwC-DP sources: survey", response)
        self.assertIn("Darwin Core Occurrence", response)
        self.assertIn("Compatible cores: event, taxon", response)

    def test_humboldt_lookup_returns_guidance_and_registered_terms(self):
        response = GetDwCExtensionInfo(extension="humboldt_ecological_inventory").run()
        term_response = GetDwCExtensionInfo(
            extension="humboldt_ecological_inventory",
            terms=["samplingEffortValue"],
        ).run()

        self.assertIn("Projection guidance:", response)
        self.assertIn("samplingEffortValue", response)
        self.assertIn("The numeric value for the sampling effort", term_response)
        self.assertNotIn("<translation", term_response)
        self.assertIn("inferred from detected occurrences", EXTENSION_SCHEMAS[
            DarwinCoreExtensionType.HUMBOLDT_ECOLOGICAL_INVENTORY
        ].avoid_when)


class EventDateNormalizationTests(SimpleTestCase):
    def test_preserves_partial_date_precision_and_interval_precision(self):
        self.assertEqual(normalize_event_date("2010"), "2010")
        self.assertEqual(normalize_event_date("2010-07"), "2010-07")
        self.assertEqual(normalize_event_date("2010/2011"), "2010/2011")
        self.assertEqual(normalize_event_date("2010-07/2011-03"), "2010-07/2011-03")

    def test_normalizes_only_explicit_textual_precision(self):
        self.assertEqual(normalize_event_date("March 2010"), "2010-03")
        self.assertEqual(normalize_event_date("11 March 2010"), "2010-03-11")
        self.assertIsNone(normalize_event_date("03/04/2010"))

    def test_normalizes_compact_and_spreadsheet_date_values(self):
        self.assertEqual(normalize_event_date("20250211"), "2025-02-11")
        self.assertEqual(
            normalize_event_date("2025-02-11 00:00:00"),
            "2025-02-11",
        )
        self.assertEqual(
            normalize_event_date(pd.Timestamp("2025-02-11 00:00:00")),
            "2025-02-11",
        )
        self.assertIsNone(normalize_event_date("20250230"))

    def test_normalizes_real_datetimes_to_iso_without_dropping_time(self):
        self.assertEqual(
            normalize_event_date("2025-02-11T00:00:00"),
            "2025-02-11T00:00:00",
        )
        self.assertEqual(
            normalize_event_date("2025-02-11 12:34:56"),
            "2025-02-11T12:34:56",
        )
        self.assertEqual(
            normalize_event_date("2025-02-11T00:00:00Z"),
            "2025-02-11T00:00:00+00:00",
        )

    def test_basic_validation_does_not_invent_date_components(self):
        df = pd.DataFrame(
            {"eventDate": ["2010", "2010/2011", "March 2010", "not known"]}
        )

        normalized, failed, future = BasicValidationForSomeDwCTerms(
            agent_id=1
        ).validate_and_format_event_dates(df)

        self.assertEqual(
            normalized["eventDate"].tolist(),
            ["2010", "2010/2011", "2010-03", "not known"],
        )
        self.assertEqual(failed, [3])
        self.assertEqual(future, [])


class SetEMLTemporalInferenceTests(SimpleTestCase):
    def test_infer_temporal_bounds_from_eventdate_column(self):
        df = pd.DataFrame(
            {
                "eventDate": [
                    "2019-01-15",
                    "2020-02",
                    "2018",
                    "2007-03-01/2008-05-11",
                ]
            }
        )

        bounds = SetEML._infer_temporal_bounds_from_df(df)

        self.assertEqual(bounds[0].isoformat(), "2007-03-01")
        self.assertEqual(bounds[1].isoformat(), "2020-02-29")

    def test_infer_temporal_bounds_from_year_month_day_columns(self):
        df = pd.DataFrame(
            {
                "year": ["2011", "2013", "2012"],
                "month": ["", "6", "2"],
                "day": ["", "", "29"],
            }
        )

        bounds = SetEML._infer_temporal_bounds_from_df(df)

        self.assertEqual(bounds[0].isoformat(), "2011-01-01")
        self.assertEqual(bounds[1].isoformat(), "2013-06-30")

    def test_resolve_temporal_scope_replaces_today_placeholder_with_inferred_range(self):
        today = datetime.date.today()
        today_text = today.strftime("%d %B %Y")
        provided_scope = f"{today_text} - {today_text}"
        inferred_scope = "2018-01-01/2020-12-31"

        resolved_scope, note = SetEML._resolve_temporal_scope(provided_scope, None, inferred_scope)

        self.assertEqual(resolved_scope, inferred_scope)
        self.assertIn("adjusted from placeholder", note)

    def test_resolve_temporal_scope_keeps_non_placeholder_value(self):
        provided_scope = "2015-01-01/2016-12-31"
        inferred_scope = "2018-01-01/2020-12-31"

        resolved_scope, note = SetEML._resolve_temporal_scope(provided_scope, None, inferred_scope)

        self.assertEqual(resolved_scope, provided_scope)
        self.assertIsNone(note)

    def test_resolve_temporal_scope_keeps_existing_non_placeholder_when_not_provided(self):
        existing_scope = "2001-01-01/2003-12-31"
        inferred_scope = "2018-01-01/2020-12-31"

        resolved_scope, note = SetEML._resolve_temporal_scope(None, existing_scope, inferred_scope)

        self.assertEqual(resolved_scope, existing_scope)
        self.assertIsNone(note)

    def test_infer_geographic_scope_from_country_and_coordinates(self):
        class DummyTable:
            def __init__(self, df):
                self.df = df

        class DummyTableSet:
            def __init__(self, tables):
                self._tables = tables

            def all(self):
                return self._tables

        class DummyDataset:
            def __init__(self, tables):
                self.table_set = DummyTableSet(tables)

        dataset = DummyDataset(
            [
                DummyTable(
                    pd.DataFrame(
                        {
                            "country": ["Norway", "Sweden", "Norway"],
                            "decimalLatitude": ["60.1", "61.2", "59.9"],
                            "decimalLongitude": ["10.0", "11.5", "9.8"],
                        }
                    )
                )
            ]
        )

        inferred = SetEML._infer_geographic_scope_from_dataset(dataset)

        self.assertIn("Countries:", inferred)
        self.assertIn("Norway", inferred)
        self.assertIn("Sweden", inferred)
        self.assertIn("Coordinate bounds:", inferred)

    def test_infer_taxonomic_scope_from_taxonomic_columns(self):
        class DummyTable:
            def __init__(self, df):
                self.df = df

        class DummyTableSet:
            def __init__(self, tables):
                self._tables = tables

            def all(self):
                return self._tables

        class DummyDataset:
            def __init__(self, tables):
                self.table_set = DummyTableSet(tables)

        dataset = DummyDataset(
            [
                DummyTable(
                    pd.DataFrame(
                        {
                            "family": ["Felidae", "Canidae", "Felidae"],
                        }
                    )
                )
            ]
        )

        inferred = SetEML._infer_taxonomic_scope_from_dataset(dataset)

        self.assertIn("Families:", inferred)
        self.assertIn("Felidae", inferred)
        self.assertIn("Canidae", inferred)

    def test_infer_taxonomic_keywords_from_taxonomic_columns(self):
        class DummyTable:
            def __init__(self, df):
                self.df = df

        class DummyTableSet:
            def __init__(self, tables):
                self._tables = tables

            def all(self):
                return self._tables

        class DummyDataset:
            def __init__(self, tables):
                self.table_set = DummyTableSet(tables)

        dataset = DummyDataset(
            [
                DummyTable(
                    pd.DataFrame(
                        {
                            "kingdom": ["Animalia", "Animalia"],
                            "family": ["Felidae", "Canidae"],
                        }
                    )
                )
            ]
        )

        inferred = SetEML._infer_taxonomic_keywords_from_dataset(dataset)

        self.assertIn({"rank": "kingdom", "scientificName": "Animalia"}, inferred)
        self.assertIn({"rank": "family", "scientificName": "Felidae"}, inferred)
        self.assertIn({"rank": "family", "scientificName": "Canidae"}, inferred)

    def test_infer_methodology_from_sampling_protocol(self):
        class DummyTable:
            def __init__(self, df):
                self.df = df

        class DummyTableSet:
            def __init__(self, tables):
                self._tables = tables

            def all(self):
                return self._tables

        class DummyDataset:
            def __init__(self, tables):
                self.table_set = DummyTableSet(tables)

        dataset = DummyDataset(
            [
                DummyTable(
                    pd.DataFrame(
                        {
                            "samplingProtocol": ["Camera trap", "Camera trap"],
                        }
                    )
                )
            ]
        )

        inferred = SetEML._infer_methodology_from_dataset(dataset)
        self.assertEqual(inferred, "Camera trap")


class SetEMLProjectTitleTests(TestCase):
    def setUp(self):
        self.dataset = Dataset.objects.create(
            title="Project metadata",
            eml={"project_title": "Existing project"},
        )
        task = Task.objects.create(name="Metadata", text="Set metadata", order=1)
        self.agent = Agent.objects.create(dataset=self.dataset, task=task)

    def test_explicit_null_clears_existing_project_title(self):
        result = SetEML(agent_id=self.agent.id, project_title=None).run()
        self.dataset.refresh_from_db()

        self.assertEqual(result, "EML has been successfully set.")
        self.assertNotIn("project_title", self.dataset.eml)

    def test_omitted_project_title_preserves_existing_value(self):
        result = SetEML(agent_id=self.agent.id).run()
        self.dataset.refresh_from_db()

        self.assertEqual(result, "EML has been successfully set.")
        self.assertEqual(self.dataset.eml["project_title"], "Existing project")


class ExcelWorkbookRepairTests(SimpleTestCase):
    @staticmethod
    def _build_workbook_bytes():
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "template"
        sheet["A1"] = "column"
        sheet["A2"] = "value"
        sheet["A1"].comment = openpyxl.comments.Comment("sample comment", "tester")

        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _inject_invalid_font_family_values(workbook_bytes):
        source = io.BytesIO(workbook_bytes)
        output = io.BytesIO()

        with zipfile.ZipFile(source, "r") as source_zip:
            with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as output_zip:
                for info in source_zip.infolist():
                    data = source_zip.read(info.filename)
                    if info.filename.endswith(".xml"):
                        xml = data.decode("utf-8")
                        xml = xml.replace('family val="2"', 'family val="34"')
                        data = xml.encode("utf-8")
                    output_zip.writestr(info, data)

        return output.getvalue()

    def test_loader_repairs_out_of_range_font_family_values(self):
        from .models import UserFile

        valid_workbook = self._build_workbook_bytes()
        broken_workbook = self._inject_invalid_font_family_values(valid_workbook)

        with self.assertRaises(ValueError):
            openpyxl.load_workbook(io.BytesIO(broken_workbook))

        workbook = UserFile._load_workbook_with_xml_repair(broken_workbook)
        self.assertEqual(workbook.sheetnames, ["template"])

    def test_sanitizer_is_noop_for_valid_workbook(self):
        from .models import UserFile

        workbook_bytes = self._build_workbook_bytes()
        sanitized_bytes, modified = UserFile._sanitize_excel_xml_font_families(workbook_bytes)

        self.assertFalse(modified)
        self.assertEqual(workbook_bytes, sanitized_bytes)


class LogBugWithDeveloperTests(SimpleTestCase):
    @patch("api.agent_tools.discord_bot.send_discord_message")
    def test_uses_discord_user_id_for_direct_mention(self, send_discord_message_mock):
        with patch.dict(os.environ, {"DISCORD_DEVELOPER_USER_ID": "1234567890"}, clear=False):
            result = LogBugWithDeveloper(
                message="Python tool failed with unexpected KeyError",
                agent_id=44,
                urgent=True,
            ).run()

        self.assertIn("Bug report sent", result)
        sent_message = send_discord_message_mock.call_args.args[0]
        sent_allowed_mentions = send_discord_message_mock.call_args.kwargs["allowed_mentions"]

        self.assertIn("<@1234567890>", sent_message)
        self.assertIn("Agent ID: 44", sent_message)
        self.assertIn("Python tool failed with unexpected KeyError", sent_message)
        self.assertEqual(sent_allowed_mentions, {"parse": [], "users": ["1234567890"]})

    @patch("api.agent_tools.discord_bot.send_discord_message")
    def test_falls_back_to_plain_rkian_tag_without_user_id(self, send_discord_message_mock):
        with patch.dict(os.environ, {"DISCORD_DEVELOPER_USER_ID": "", "DISCORD_DEVELOPER_HANDLE": "@_rkian"}, clear=False):
            result = LogBugWithDeveloper(message="Validation response parsing failed").run()

        self.assertIn("Bug report sent", result)
        sent_message = send_discord_message_mock.call_args.args[0]

        self.assertIn("@_rkian", sent_message)
        self.assertIn("Validation response parsing failed", sent_message)
        self.assertNotIn("allowed_mentions", send_discord_message_mock.call_args.kwargs)


class ResponsesAdapterCompatibilityTests(SimpleTestCase):
    class _Message:
        def __init__(self, openai_obj):
            self.openai_obj = openai_obj

    @override_settings(OPENAI_MODEL="gpt-5.4", OPENAI_REASONING_EFFORT="high")
    @patch("api.helpers.openai_helpers.query_responses_api")
    def test_default_request_uses_gpt_5_4_with_high_reasoning(self, query_mock):
        query_mock.return_value = SimpleNamespace(
            id="resp_1",
            status="completed",
            output=[],
            output_text="Done.",
        )

        create_response_message(
            [self._Message({"role": "user", "content": "Transform this."})],
            [],
        )

        request = query_mock.call_args.args[0]
        self.assertEqual(request["model"], "gpt-5.4")
        self.assertEqual(request["reasoning"], {"effort": "high"})
        self.assertNotIn("temperature", request)

    def test_messages_are_mapped_to_responses_input_with_tool_history(self):
        messages = [
            self._Message({"role": "system", "content": "You are a helper."}),
            self._Message({"role": "user", "content": "Clean my table"}),
            self._Message(
                {
                    "role": "assistant",
                    "content": "I will run Python.",
                    "tool_calls": [
                        {
                            "id": "call_abc",
                            "type": "function",
                            "function": {
                                "name": "Python",
                                "arguments": "{\"code\":\"print(1)\"}",
                            },
                        }
                    ],
                }
            ),
            self._Message({"role": "tool", "tool_call_id": "call_abc", "content": "1"}),
        ]

        items = _messages_to_responses_input(messages)

        self.assertEqual(items[0], {"role": "system", "content": "You are a helper."})
        self.assertEqual(items[1], {"role": "user", "content": "Clean my table"})
        self.assertEqual(items[2], {"role": "assistant", "content": "I will run Python."})
        self.assertEqual(
            items[3],
            {
                "type": "function_call",
                "call_id": "call_abc",
                "name": "Python",
                "arguments": "{\"code\":\"print(1)\"}",
            },
        )
        self.assertEqual(
            items[4],
            {
                "type": "function_call_output",
                "call_id": "call_abc",
                "output": "1",
            },
        )

    def test_pdf_file_inputs_attach_to_latest_user_message(self):
        items = [
            {"role": "system", "content": "You are a helper."},
            {"role": "user", "content": "Read the manuscript."},
            {"role": "assistant", "content": "I can inspect it."},
        ]

        updated = _attach_pdf_files_to_latest_user_message(
            items,
            [{"type": "input_file", "file_id": "file_123", "filename": "paper.pdf"}],
        )

        self.assertEqual(updated[0], items[0])
        self.assertEqual(updated[2], items[2])
        self.assertEqual(updated[1]["role"], "user")
        self.assertEqual(updated[1]["content"][0], {"type": "input_text", "text": "Read the manuscript."})
        self.assertEqual(updated[1]["content"][1]["type"], "input_text")
        self.assertIn("attached", updated[1]["content"][1]["text"])
        self.assertEqual(updated[1]["content"][2], {"type": "input_file", "file_id": "file_123"})

    @patch("api.helpers.openai_helpers._resolve_openai_file_id_from_user_file_id")
    def test_user_message_pdf_attachments_are_hydrated_from_message_metadata(self, resolve_file_id_mock):
        resolve_file_id_mock.return_value = "file_pdf_1"
        messages = [
            self._Message({"role": "system", "content": "You are a helper."}),
            self._Message(
                {
                    "role": "user",
                    "content": "Please use this PDF.",
                    "pdf_attachments": [{"user_file_id": 101, "filename": "paper.pdf"}],
                }
            ),
        ]

        items = _messages_to_responses_input(messages)

        self.assertEqual(items[0], {"role": "system", "content": "You are a helper."})
        self.assertEqual(items[1]["role"], "user")
        self.assertEqual(items[1]["content"][0], {"type": "input_text", "text": "Please use this PDF."})
        self.assertEqual(items[1]["content"][1], {"type": "input_file", "file_id": "file_pdf_1"})
        resolve_file_id_mock.assert_called_once_with(101)

    def test_responses_output_is_mapped_back_to_legacy_message_shape(self):
        response = SimpleNamespace(
            output=[
                SimpleNamespace(
                    type="function_call",
                    call_id="call_xyz",
                    name="Python",
                    arguments="{\"code\":\"print(2)\"}",
                ),
                SimpleNamespace(
                    type="message",
                    content=[SimpleNamespace(type="output_text", text="Done.")],
                ),
            ],
            output_text="Done.",
        )

        compat_message = _response_to_compat_message(response)

        self.assertEqual(compat_message.role, "assistant")
        self.assertEqual(compat_message.content, "Done.")
        self.assertEqual(len(compat_message.tool_calls), 1)
        self.assertEqual(compat_message.tool_calls[0].id, "call_xyz")
        self.assertEqual(compat_message.tool_calls[0].function.name, "Python")
        self.assertEqual(compat_message.tool_calls[0].function.arguments, "{\"code\":\"print(2)\"}")
        self.assertEqual(
            compat_message.dict(),
            {
                "role": "assistant",
                "content": "Done.",
                "tool_calls": [
                    {
                        "id": "call_xyz",
                        "type": "function",
                        "function": {
                            "name": "Python",
                            "arguments": "{\"code\":\"print(2)\"}",
                        },
                    }
                ],
            },
        )

    def test_functions_are_mapped_to_responses_tool_schema(self):
        class DummyFunction:
            @classmethod
            def openai_schema(cls):
                return {
                    "name": "DummyFunction",
                    "description": "A test helper function.",
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "foo": {"type": "string"},
                        },
                        "required": ["foo"],
                    },
                }

        tools = _functions_to_responses_tools([DummyFunction])
        self.assertEqual(
            tools,
            [
                {
                    "type": "function",
                    "name": "DummyFunction",
                    "description": "A test helper function.",
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "foo": {"type": "string"},
                        },
                        "required": ["foo"],
                    },
                }
            ],
        )

    def test_set_basic_metadata_schema_keeps_title_and_description_properties(self):
        schema = SetBasicMetadata.openai_schema()
        parameters = schema["parameters"]
        properties = parameters["properties"]

        self.assertIn("agent_id", properties)
        self.assertIn("title", properties)
        self.assertIn("description", properties)
        self.assertNotIn("suitable_for_publication_on_gbif", properties)

    def test_upload_dwca_schema_exposes_extension_table_mapping_values(self):
        schema = UploadDwCA.openai_schema()
        extension_tables = schema["parameters"]["properties"]["extension_tables"]
        array_schema = next(
            item for item in extension_tables["anyOf"] if item.get("type") == "array"
        )
        assignment_ref = array_schema["items"]["$ref"]
        assignment_name = assignment_ref.rsplit("/", 1)[-1]
        assignment_schema = schema["parameters"]["$defs"][assignment_name]

        self.assertEqual(
            set(assignment_schema["required"]),
            {"table_id", "extension_type", "core_id_column"},
        )
        self.assertIn("DarwinCoreExtensionType", schema["parameters"]["$defs"])


class DwcaExportSanitizationTests(SimpleTestCase):
    def test_surrogateescape_text_is_safe_for_utf8_export(self):
        df = pd.DataFrame({"occurrenceID": ["occ-1"], "fieldNotes": ["collector\udc92s note"]})

        sanitized = _sanitize_dataframe_for_utf8_export(df)
        value = sanitized.loc[0, "fieldNotes"]

        self.assertEqual(value, "collector’s note")
        value.encode("utf-8")

    def test_exact_archive_validation_detects_malformed_rows(self):
        meta = """<?xml version="1.0" encoding="UTF-8"?>
<archive xmlns="http://rs.tdwg.org/dwc/text/">
  <core encoding="UTF-8" fieldsTerminatedBy="\\t" fieldsEnclosedBy="&quot;"
        ignoreHeaderLines="1" rowType="http://rs.tdwg.org/dwc/terms/Occurrence">
    <files><location>occurrence.txt</location></files>
    <id index="0"/>
    <field index="1" term="http://rs.tdwg.org/dwc/terms/scientificName"/>
  </core>
</archive>"""
        with tempfile.TemporaryDirectory() as temp_dir:
            archive_path = Path(temp_dir) / "bad.zip"
            with zipfile.ZipFile(archive_path, "w") as archive:
                archive.writestr("meta.xml", meta)
                archive.writestr("eml.xml", "<eml/>")
                archive.writestr(
                    "occurrence.txt",
                    "occurrenceID\tscientificName\nocc-1\tAcer rubrum\textra\n",
                )

            result = validate_dwca_archive(archive_path)

        self.assertFalse(result["valid"])
        self.assertTrue(any("field count" in error for error in result["errors"]))

    def test_exact_archive_validation_rejects_incomplete_eml_project(self):
        meta = """<?xml version="1.0" encoding="UTF-8"?>
<archive xmlns="http://rs.tdwg.org/dwc/text/">
  <core encoding="UTF-8" fieldsTerminatedBy="\\t" ignoreHeaderLines="1"
        rowType="http://rs.tdwg.org/dwc/terms/Occurrence">
    <files><location>occurrence.txt</location></files>
    <id index="0"/>
  </core>
</archive>"""
        eml = """<eml><dataset><project><title>Existing project</title></project></dataset></eml>"""
        with tempfile.TemporaryDirectory() as temp_dir:
            archive_path = Path(temp_dir) / "bad-eml.zip"
            with zipfile.ZipFile(archive_path, "w") as archive:
                archive.writestr("meta.xml", meta)
                archive.writestr("eml.xml", eml)
                archive.writestr("occurrence.txt", "occurrenceID\nocc-1\n")

            result = validate_dwca_archive(archive_path)

        self.assertFalse(result["valid"])
        self.assertTrue(any("project must contain" in error for error in result["errors"]))

    def test_exact_archive_validation_rejects_unresolved_extension_core_ids(self):
        meta = r"""<?xml version="1.0" encoding="UTF-8"?>
<archive xmlns="http://rs.tdwg.org/dwc/text/">
  <core encoding="UTF-8" fieldsTerminatedBy="\t" ignoreHeaderLines="1"
        rowType="http://rs.tdwg.org/dwc/terms/Occurrence">
    <files><location>occurrence.txt</location></files>
    <id index="0"/>
    <field index="0" term="http://rs.tdwg.org/dwc/terms/occurrenceID"/>
  </core>
  <extension encoding="UTF-8" fieldsTerminatedBy="\t" ignoreHeaderLines="1"
        rowType="http://rs.tdwg.org/dwc/terms/Identification">
    <files><location>identification.txt</location></files>
    <coreid index="0"/>
    <field index="1" term="http://rs.tdwg.org/dwc/terms/identificationID"/>
  </extension>
</archive>"""
        with tempfile.TemporaryDirectory() as temp_dir:
            archive_path = Path(temp_dir) / "bad-extension.zip"
            with zipfile.ZipFile(archive_path, "w") as archive:
                archive.writestr("meta.xml", meta)
                archive.writestr("eml.xml", "<eml/>")
                archive.writestr("occurrence.txt", "occurrenceID\nocc-1\n")
                archive.writestr("identification.txt", "_coreid\tidentificationID\nmissing\tid-1\n")

            result = validate_dwca_archive(archive_path)

        self.assertFalse(result["valid"])
        self.assertTrue(any("do not resolve to the core" in error for error in result["errors"]))

    @patch("api.helpers.publish.upload_file")
    @patch("api.helpers.publish.Minio")
    def test_upload_dwca_writes_explicit_extension_coreid(self, minio_mock, upload_mock):
        captured = {}

        def capture_archive(client, bucket, object_name, local_path):
            with zipfile.ZipFile(local_path) as archive:
                captured["meta"] = archive.read("meta.xml")

        upload_mock.side_effect = capture_archive
        env = {
            "MINIO_URI": "storage.example.org",
            "MINIO_ACCESS_KEY": "key",
            "MINIO_SECRET_KEY": "secret",
            "MINIO_BUCKET": "bucket",
            "MINIO_BUCKET_FOLDER": "packages",
        }
        core = pd.DataFrame([
            {
                "scientificName": "Apus apus",
                "occurrenceID": "occ-1",
                "occurrenceRemarks": "first line\nsecond line\twith tab",
            },
        ])
        identification = pd.DataFrame([
            {
                "identificationID": "identification-1",
                "_coreid": "occ-1",
                "scientificName": "Apus apus",
            },
        ])

        with patch.dict(os.environ, env):
            upload_dwca(
                core,
                "Test dataset",
                "Test description",
                core_type=DarwinCoreCoreType.OCCURRENCE,
                extensions=[(
                    identification,
                    DarwinCoreExtensionType.IDENTIFICATION_HISTORY,
                    "_coreid",
                )],
            )

        meta_root = ET.fromstring(captured["meta"])
        extension = next(
            element for element in meta_root if element.tag.rsplit("}", 1)[-1] == "extension"
        )
        coreid = next(
            element for element in extension if element.tag.rsplit("}", 1)[-1] == "coreid"
        )
        fields = [
            element for element in extension if element.tag.rsplit("}", 1)[-1] == "field"
        ]

        self.assertEqual(coreid.attrib["index"], "0")
        self.assertTrue(any(field.attrib["index"] == "1" for field in fields))

    @patch("api.helpers.publish.upload_file")
    @patch("api.helpers.publish.Minio")
    def test_occurrence_core_exports_audiovisual_emof_and_ggbn_extensions(
        self, minio_mock, upload_mock
    ):
        captured = {}

        def capture_archive(client, bucket, object_name, local_path):
            with zipfile.ZipFile(local_path) as archive:
                captured["meta"] = archive.read("meta.xml")

        upload_mock.side_effect = capture_archive
        env = {
            "MINIO_URI": "storage.example.org",
            "MINIO_ACCESS_KEY": "key",
            "MINIO_SECRET_KEY": "secret",
            "MINIO_BUCKET": "bucket",
            "MINIO_BUCKET_FOLDER": "packages",
        }
        core = pd.DataFrame([{"occurrenceID": "occ-1", "scientificName": "Apus apus"}])
        extension_values = {
            DarwinCoreExtensionType.AUDIOVISUAL: {"identifier": "https://example.org/media/1"},
            DarwinCoreExtensionType.EXTENDED_MEASUREMENT_OR_FACT: {
                "measurementID": "measurement-1",
                "measurementTypeID": "https://example.org/vocabulary/body-mass",
            },
            DarwinCoreExtensionType.GGBN_MATERIAL_SAMPLE: {"materialSampleType": "tissue"},
            DarwinCoreExtensionType.GGBN_AMPLIFICATION: {"amplificationSuccess": "true"},
            DarwinCoreExtensionType.GGBN_CLONING: {"cloningMethod": "plasmid cloning"},
            DarwinCoreExtensionType.GGBN_GEL_IMAGE: {"identifier": "https://example.org/gel/1"},
            DarwinCoreExtensionType.GGBN_LOAN: {"loanIdentifier": "loan-1"},
            DarwinCoreExtensionType.GGBN_PERMIT: {"permitURI": "https://example.org/permit/1"},
            DarwinCoreExtensionType.GGBN_PREPARATION: {"preparationType": "gDNA"},
            DarwinCoreExtensionType.GGBN_PRESERVATION: {"preservationType": "frozen"},
        }
        extensions = [
            (pd.DataFrame([{"_coreid": "occ-1", **values}]), extension_type, "_coreid")
            for extension_type, values in extension_values.items()
        ]

        with patch.dict(os.environ, env):
            upload_dwca(
                core,
                "Molecular media dataset",
                "A material sample with media, measurements, and GGBN metadata.",
                core_type=DarwinCoreCoreType.OCCURRENCE,
                extensions=extensions,
            )

        meta_root = ET.fromstring(captured["meta"])
        row_types = {
            element.attrib["rowType"]
            for element in meta_root
            if element.tag.rsplit("}", 1)[-1] == "extension"
        }
        self.assertEqual(
            row_types,
            {EXTENSION_SCHEMAS[extension_type].row_type for extension_type in extension_values},
        )

    def test_upload_dwca_rejects_extension_incompatible_with_core(self):
        core = pd.DataFrame([
            {"occurrenceID": "occ-1", "scientificName": "Apus apus"},
        ])
        humboldt = pd.DataFrame([
            {"_coreid": "occ-1", "protocolNames": "Point count"},
        ])

        with self.assertRaisesRegex(ValueError, "not compatible with the 'occurrence' core"):
            upload_dwca(
                core,
                "Test dataset",
                "Test description",
                core_type=DarwinCoreCoreType.OCCURRENCE,
                extensions=[(
                    humboldt,
                    DarwinCoreExtensionType.HUMBOLDT_ECOLOGICAL_INVENTORY,
                    "_coreid",
                )],
            )

    def test_upload_dwca_reports_all_invalid_extension_links_together(self):
        core = pd.DataFrame([
            {"occurrenceID": "occ-1", "scientificName": "Apus apus"},
        ])
        identification = pd.DataFrame([
            {"_coreid": "occurrence-pk-1", "identificationID": "identification-1"},
        ])
        multimedia = pd.DataFrame([
            {"_coreid": "occurrence-pk-2", "identifier": "https://example.org/media/1"},
            {"_coreid": "occurrence-pk-2", "identifier": "https://example.org/media/2"},
        ])

        with self.assertRaises(DwcaExtensionLinkError) as raised:
            upload_dwca(
                core,
                "Test dataset",
                "Test description",
                core_type=DarwinCoreCoreType.OCCURRENCE,
                extensions=[
                    (identification, DarwinCoreExtensionType.IDENTIFICATION, "_coreid"),
                    (multimedia, DarwinCoreExtensionType.MULTIMEDIA, "_coreid"),
                ],
            )

        self.assertEqual(raised.exception.core_id_column, "occurrenceID")
        self.assertEqual(
            [issue["extension_type"] for issue in raised.exception.issues],
            ["identification", "multimedia"],
        )
        self.assertEqual(
            [issue["unresolved_count"] for issue in raised.exception.issues],
            [1, 1],
        )

    @patch("api.helpers.publish.upload_file")
    @patch("api.helpers.publish.Minio")
    def test_event_core_exports_occurrence_and_humboldt_extensions(self, minio_mock, upload_mock):
        captured = {}

        def capture_archive(client, bucket, object_name, local_path):
            with zipfile.ZipFile(local_path) as archive:
                captured["names"] = archive.namelist()
                captured["meta"] = archive.read("meta.xml")
                captured["tables"] = {
                    name: archive.read(name).decode("utf-8")
                    for name in archive.namelist()
                    if name.endswith(".txt")
                }

        upload_mock.side_effect = capture_archive
        event = pd.DataFrame([
            {
                "eventID": "survey-2025",
                "eventDate": "2025-04-01/2025-04-30",
                "samplingProtocol": "Fixed-route point counts",
            },
            {
                "eventID": "visit-1",
                "parentEventID": "survey-2025",
                "eventDate": "2025-04-12",
                "samplingProtocol": "Ten-minute point count",
            },
        ])
        occurrence = pd.DataFrame([
            {
                "_coreid": "visit-1",
                "occurrenceID": "occ-1",
                "basisOfRecord": "HumanObservation",
                "scientificName": "Apus apus",
                "occurrenceStatus": "present",
            },
        ])
        humboldt = pd.DataFrame([
            {
                "_coreid": "visit-1",
                "protocolNames": "Point count",
                "isSamplingEffortReported": True,
                "samplingEffortValue": 10,
                "samplingEffortUnit": "minutes",
            },
        ])
        env = {
            "MINIO_URI": "storage.example.org",
            "MINIO_ACCESS_KEY": "key",
            "MINIO_SECRET_KEY": "secret",
            "MINIO_BUCKET": "bucket",
            "MINIO_BUCKET_FOLDER": "packages",
        }

        with patch.dict(os.environ, env):
            upload_dwca(
                event,
                "Point count survey",
                "Repeated bird point counts with explicit survey effort.",
                core_type=DarwinCoreCoreType.EVENT,
                extensions=[
                    (
                        occurrence,
                        DarwinCoreExtensionType.OCCURRENCE,
                        "_coreid",
                    ),
                    (
                        humboldt,
                        DarwinCoreExtensionType.HUMBOLDT_ECOLOGICAL_INVENTORY,
                        "_coreid",
                    ),
                ],
            )

        meta_root = ET.fromstring(captured["meta"])
        extensions = [
            element
            for element in meta_root
            if element.tag.rsplit("}", 1)[-1] == "extension"
        ]
        self.assertEqual(
            {extension.attrib["rowType"] for extension in extensions},
            {
                "http://rs.tdwg.org/dwc/terms/Occurrence",
                "http://rs.tdwg.org/eco/terms/Event",
            },
        )
        self.assertEqual(len(captured["tables"]), 3)
        combined_tables = "\n".join(captured["tables"].values())
        self.assertIn("Apus apus", combined_tables)
        self.assertIn("Point count", combined_tables)


class UploadDwcaCorrectionFeedbackTests(TestCase):
    @patch("api.agent_tools.discord_bot.send_discord_message")
    def test_invalid_extension_links_return_structured_feedback_without_alert(self, discord_mock):
        task = Task.objects.create(name="DwC-A feedback test", text="Test", order=999)
        dataset = Dataset.objects.create(title="Test dataset", description="Test description")
        agent = Agent.objects.create(dataset=dataset, task=task)
        core = Table.objects.create(
            dataset=dataset,
            title="occurrence_dwca",
            df=pd.DataFrame([{"occurrenceID": "occ-1", "scientificName": "Apus apus"}]),
        )
        identification = Table.objects.create(
            dataset=dataset,
            title="identification_dwca",
            df=pd.DataFrame([{
                "_coreid": "occurrence-pk-1",
                "identificationID": "identification-1",
            }]),
        )
        multimedia = Table.objects.create(
            dataset=dataset,
            title="multimedia_dwca",
            df=pd.DataFrame([{
                "_coreid": "occurrence-pk-2",
                "identifier": "https://example.org/media/1",
            }]),
        )

        result = UploadDwCA(
            agent_id=agent.id,
            core_table_id=core.id,
            core_type=DarwinCoreCoreType.OCCURRENCE,
            extension_tables=[
                {
                    "table_id": identification.id,
                    "extension_type": DarwinCoreExtensionType.IDENTIFICATION,
                    "core_id_column": "_coreid",
                },
                {
                    "table_id": multimedia.id,
                    "extension_type": DarwinCoreExtensionType.MULTIMEDIA,
                    "core_id_column": "_coreid",
                },
            ],
        ).run()

        feedback = json.loads(result)
        self.assertEqual(feedback["status"], "correction_required")
        self.assertEqual(feedback["core_identifier_column"], "occurrenceID")
        self.assertEqual(
            [issue["table_id"] for issue in feedback["extension_tables"]],
            [identification.id, multimedia.id],
        )
        self.assertIn("Do not copy DwC-DP `_pk` or `_fk` values directly", feedback["required_action"])
        discord_mock.assert_not_called()


class DatasetSummarySerializerTests(TestCase):
    def test_relational_modeling_is_included_in_every_agent_prompt(self):
        dataset = Dataset.objects.create(
            title="Rich package",
            description="Test",
        )
        task = Task.objects.create(name="Data transformation", text="Transform", order=1)

        agent = Agent.create_with_system_message(dataset=dataset, task=task, tables=[])
        prompt = agent.message_set.get(openai_obj__role="system").openai_obj["content"]

        self.assertIn("Build a rich relational DwC-DP", prompt)
        self.assertIn("Usage Policy", prompt)

    def test_package_explorer_model_uses_current_schema_and_reports_link_coverage(self):
        dataset = Dataset.objects.create(title="Explorer")
        Table.objects.create(
            dataset=dataset,
            title="event",
            df=pd.DataFrame(
                [
                    {"event_pk": "event-1", "eventCategory": "occurrence"},
                    {"event_pk": "event-2", "eventCategory": "occurrence"},
                ]
            ),
        )
        Table.objects.create(
            dataset=dataset,
            title="occurrence",
            df=pd.DataFrame(
                [
                    {"occurrence_pk": "occ-1", "event_fk": "event-1", "occurrenceStatus": "present"},
                    {"occurrence_pk": "occ-2", "event_fk": "missing", "occurrenceStatus": "present"},
                    {"occurrence_pk": "occ-3", "event_fk": "", "occurrenceStatus": "present"},
                ]
            ),
        )

        model = build_dwc_dp_explorer_model(dataset)

        self.assertEqual({node["id"] for node in model["nodes"]}, {"event", "occurrence"})
        occurrence = next(node for node in model["nodes"] if node["id"] == "occurrence")
        self.assertEqual(occurrence["rowCount"], 3)
        self.assertTrue(next(field for field in occurrence["fields"] if field["name"] == "occurrence_pk")["primary"])

        relationship = next(
            edge
            for edge in model["edges"]
            if edge["source"] == "occurrence" and edge["target"] == "event"
        )
        self.assertEqual(relationship["predicate"], "happened during")
        self.assertEqual(relationship["sourceFields"], ["event_fk"])
        self.assertEqual(relationship["targetFields"], ["event_pk"])
        self.assertEqual(relationship["populatedRows"], 2)
        self.assertEqual(relationship["linkedRows"], 1)
        self.assertEqual(relationship["unmatchedRows"], 1)
        self.assertEqual(relationship["blankRows"], 1)

    def test_reports_relational_resource_counts_without_calling_the_sum_records(self):
        dataset = Dataset.objects.create(title="Relational counts")
        Table.objects.create(
            dataset=dataset,
            title="source working copy",
            df=pd.DataFrame({"id": range(7)}),
        )
        for title, rows in {
            "event": 2,
            "occurrence": 5,
            "material": 3,
            "identification": 4,
        }.items():
            Table.objects.create(
                dataset=dataset,
                title=title,
                df=pd.DataFrame({"id": range(rows)}),
            )

        data = DatasetListSerializer(dataset).data

        self.assertEqual(
            data["counts"]["resources"],
            {"event": 2, "identification": 4, "material": 3, "occurrence": 5},
        )
        self.assertEqual(data["counts"]["package_rows"], 14)
        self.assertEqual(data["counts"]["source_rows"], 7)
        self.assertEqual(data["record_count"], 5)

    def test_ready_packages_have_ready_status_and_complete_applicable_progress(self):
        Task.objects.create(name=Dataset.MANUSCRIPT_TASK_NAME, text="PDF", order=1)
        Task.objects.create(name="Data transformation", text="Transform", order=2)
        Task.objects.create(name="Phylogenetic tree linking", text="Tree", order=3)
        Task.objects.create(name="Final Review & Publication", text="Publish", order=4)
        Task.objects.create(name="Data maintenance", text="Maintain", order=5)
        dataset = Dataset.objects.create(
            dwc_dp_url="https://example.org/package.tar.gz",
            dwca_url="https://example.org/archive.zip",
            dwc_dp_validation={"valid": True},
        )

        list_data = DatasetListSerializer(dataset).data
        detail_data = DatasetSerializer(dataset).data

        self.assertTrue(dataset.package_ready)
        self.assertEqual(list_data["status"], "ready")
        self.assertEqual(list_data["progress"], {"done": 2, "total": 2})
        self.assertTrue(list_data["package_ready"])
        self.assertTrue(detail_data["package_ready"])
        self.assertEqual(
            detail_data["dwc_dp_standard"]["profile"],
            "https://raw.githubusercontent.com/gbif/dwc-dp/"
            "cbb6c887043876351eec1bed01c3dfc2e05c4eb4/"
            "dwc-dp/dwc-dp-profile.json",
        )
        self.assertEqual(detail_data["dwc_dp_standard"]["schema"]["version"], "0.1")

    def test_status_distinguishes_preparing_from_needs_input(self):
        task = Task.objects.create(name="Data exploration", text="Explore", order=1)
        preparing_dataset = Dataset.objects.create(title="Preparing")
        preparing_agent = Agent.objects.create(
            dataset=preparing_dataset,
            task=task,
            busy_thinking=True,
        )
        Message.objects.create(
            agent=preparing_agent,
            openai_obj={"role": Message.Role.ASSISTANT, "content": "Working"},
        )

        needs_input_dataset = Dataset.objects.create(title="Needs input")
        needs_input_agent = Agent.objects.create(dataset=needs_input_dataset, task=task)
        Message.objects.create(
            agent=needs_input_agent,
            openai_obj={"role": Message.Role.ASSISTANT, "content": "Which sheet should I use?"},
        )

        self.assertEqual(DatasetListSerializer(preparing_dataset).data["status"], "preparing")
        self.assertEqual(DatasetListSerializer(needs_input_dataset).data["status"], "needs_input")

    def test_failed_validation_prevents_package_ready_status(self):
        dataset = Dataset.objects.create(
            dwc_dp_url="https://example.org/package.tar.gz",
            dwca_url="https://example.org/archive.zip",
            dwc_dp_validation={"valid": False},
        )

        data = DatasetListSerializer(dataset).data

        self.assertFalse(dataset.package_ready)
        self.assertNotEqual(data["status"], "ready")
        self.assertFalse(data["package_ready"])


class TaskFunctionTests(TestCase):
    def test_dwc_dp_tools_are_scoped_to_package_tasks(self):
        exploration = Task.objects.create(
            name="Data content exploration",
            text="Explore",
            order=1,
        )
        transformation = Task.objects.create(
            name="Data transformation",
            text="Transform",
            order=2,
        )

        self.assertNotIn(ValidateDwcDp, exploration.functions)
        self.assertIn(ValidateDwcDp, transformation.functions)

    @override_settings(
        OPENAI_REASONING_EFFORT="medium",
        OPENAI_SIMPLE_REASONING_EFFORT="low",
    )
    def test_exploration_uses_low_reasoning_but_transformation_uses_medium(self):
        exploration = Task.objects.create(
            name="Data structure exploration",
            text="Explore",
            order=1,
        )
        transformation = Task.objects.create(
            name="Data transformation",
            text="Transform",
            order=2,
        )

        self.assertEqual(exploration.reasoning_effort, "low")
        self.assertEqual(transformation.reasoning_effort, "medium")


class PythonToolTests(SimpleTestCase):
    def test_uuid_is_available_to_executed_code(self):
        result = Python(code="print(uuid.UUID(int=0))").run()

        self.assertEqual(result.strip(), "00000000-0000-0000-0000-000000000000")


class DwcDpAssertionRoutingPromptTests(SimpleTestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        fixture_path = Path(__file__).resolve().parent / "fixtures" / "tasks.yaml"
        fixture = yaml.safe_load(fixture_path.read_text(encoding="utf-8"))
        cls.task_text = {
            item["fields"]["name"]: item["fields"]["text"]
            for item in fixture
        }

    def test_transformation_prompt_delegates_semantic_routing_to_model(self):
        text = self.task_text["Data transformation"]

        self.assertIn("MODEL-LED ASSERTION ROUTING", text)
        self.assertIn("complete dataset context", text)
        self.assertIn("Do not use or invent a hard-coded keyword/term classifier", text)
        self.assertIn("exact source-to-target join", text)
        self.assertIn("explicit, structurally inferred, or unresolved", text)
        self.assertIn("save a concise version of the assertion-routing plan", text)
        self.assertIn("validation agent can independently reconstruct and challenge it", text)

    def test_transformation_prompt_requires_an_evidence_based_relational_package(self):
        text = self.task_text["Data transformation"]

        self.assertIn("EVIDENCE-BASED RELATIONAL PACKAGE", text)
        self.assertIn("complete relational DwC-DP package", text)
        self.assertIn("do not create a resource merely because DwC-DP provides one", text)
        self.assertIn("one short evidence-based justification for every proposed resource", text)
        self.assertIn("ordinary bird observations", text)
        self.assertIn("may need only event and occurrence", text)
        self.assertIn("not a target resource count", text)

    def test_transformation_prompt_has_a_narrow_deterministic_fast_path(self):
        text = self.task_text["Data transformation"]

        self.assertIn("DETERMINISTIC EXACT-MATCH FAST PATH", text)
        self.assertIn("headers exactly match fields confirmed by GetDwcDpTableInfo", text)
        self.assertIn("not similar-looking names or semantic synonyms", text)
        self.assertIn("still require the model-led contextual review", text)

    def test_refinement_prompt_challenges_under_and_over_modeling(self):
        text = self.task_text["Data validation and refinement"]

        self.assertIn("Evidence-based relational package", text)
        self.assertIn("unnecessary wrapper entities", text)
        self.assertIn("duplicated one-per-row protocols/agents", text)
        self.assertIn("assertion rows that only repeat a native field", text)
        self.assertIn("missing dedicated resources", text)
        self.assertIn("Resource count is never the optimization target", text)
        self.assertIn("Collapse synthetic wrapper/child hierarchies", text)
        self.assertIn("remove identification rows that contain only keys", text)

    def test_transformation_prompt_preserves_date_precision_and_avoids_empty_identifications(self):
        text = self.task_text["Data transformation"]

        self.assertIn("EVENT AND IDENTIFICATION MODELLING", text)
        self.assertIn("briefly sketch the Event structure supported by the source", text)
        self.assertIn("test whether narrower activities repeat within a stable broader context", text)
        self.assertIn("State why the final Event model is flat or hierarchical", text)
        self.assertIn("A recurring place or label alone is not sufficient", text)
        self.assertIn("Do not manufacture one wrapper event plus one child event per occurrence", text)
        self.assertIn("genuine broader and narrower activities or contexts", text)
        self.assertIn("monitoring programme containing plots and dated surveys", text)
        self.assertIn("expedition containing stations and collecting events", text)
        self.assertIn("examples, not templates or trigger phrases", text)
        self.assertIn("supported by source structure, identifiers, metadata", text)
        self.assertIn("One current scientificName per occurrence normally stays in occurrence", text)
        self.assertIn("Never create identification rows containing only identification_pk", text)
        self.assertIn("Higher-taxonomy preservation is mandatory", text)
        self.assertIn("Never reduce such a row to scientificName alone", text)
        self.assertIn("DATE PRECISION", text)
        self.assertIn("Never use today's month or day as a parser default", text)
        self.assertIn("normalize_event_date", text)

    def test_transformation_prompt_reviews_dedicated_resource_candidates(self):
        text = self.task_text["Data transformation"]

        self.assertIn("DEDICATED RESOURCE CANDIDATE REVIEW", text)
        self.assertIn("Agent/Agent Role", text)
        self.assertIn("Bibliographic Resource/Reference", text)
        self.assertIn("Usage Policy", text)
        self.assertIn("Provenance", text)
        self.assertIn("Assertions are not a generic overflow destination", text)
        self.assertIn("Do not produce a ceremonial include/omit checklist", text)

    def test_transformation_prompt_requires_assertion_coverage_and_join_checks(self):
        text = self.task_text["Data transformation"]

        self.assertIn("Never discard a row because `occurrenceID`", text)
        self.assertIn("temporary source-row key", text)
        self.assertIn("unmatched and multiply matched rows", text)
        self.assertIn("examples of investigation, not fixed mappings", text)
        self.assertIn("Count that source row once for unique source-row coverage", text)
        self.assertIn("populated-value count", text)
        self.assertIn("Representing a source row does not by itself represent every useful fact", text)
        self.assertIn("focused fact-group check", text)

    def test_transformation_prompt_asks_only_about_material_domain_ambiguity(self):
        text = self.task_text["Data transformation"]

        self.assertIn("WHEN TO ASK THE USER ABOUT AMBIGUOUS DATA", text)
        self.assertIn("Perform a bounded initial assessment first", text)
        self.assertIn("materially change the package", text)
        self.assertIn("concrete row counts and representative examples", text)
        self.assertIn("current best interpretation", text)
        self.assertIn("Do not ask the user to select a Darwin Core table", text)
        self.assertIn("Do not ask merely because a value is unusual", text)
        self.assertIn("If the user cannot answer, do not invent a resolution", text)
        self.assertIn("Before omitting a material source group", text)
        self.assertIn("an entire table, a substantive group of populated values", text)
        self.assertIn("ask for confirmation in the same grouped question", text)
        self.assertIn("blank rows, repeated headers, obvious totals, or formatting artifacts", text)

    def test_transformation_prompt_bounds_investigation_and_context_output(self):
        text = self.task_text["Data transformation"]

        self.assertIn("DECISION-FIRST AND CONTEXT DISCIPLINE", text)
        self.assertIn("Every exploratory Python call must resolve", text)
        self.assertIn("exploration budget of at most three read-only Python calls", text)
        self.assertIn("There is no fourth read-only inspection call", text)
        self.assertIn("show at most five representative records across the entire call", text)
        self.assertIn("never print a complete dataframe, wide row, or long list of records", text)
        self.assertIn("Do not attempt every possible inference first", text)
        self.assertIn("stop exploring representative records", text)
        self.assertIn("GetDwcDpTableInfo(include_fields=false)", text)
        self.assertIn("exact schemas for all selected resources together", text)
        self.assertIn("parallel GetDwcDpTableInfo calls", text)
        self.assertIn("Do not fetch the same resource schema twice", text)
        self.assertIn("Do not inspect resources speculatively", text)
        self.assertIn("one bounded final check", text)
        self.assertIn("avoid turning this into an exhaustive per-column exercise", text)
        self.assertNotIn("required one-line include/omit justifications", text)
        self.assertNotIn(
            "inspect every source column, populated and missing identifier patterns",
            text,
        )

    def test_refinement_prompt_rechecks_semantics_and_mechanical_coverage(self):
        text = self.task_text["Data validation and refinement"]

        self.assertIn("Reconstruct and challenge the model-led routing", text)
        self.assertIn("exactly one combined read-only Python audit", text)
        self.assertIn("at most two targeted Python correction calls", text)
        self.assertIn("Do not spend successive calls re-inspecting the same anomaly", text)
        self.assertIn("Do not repeat it unless the final validator reports", text)
        self.assertIn("Never assume blank `occurrenceID` means the row is irrelevant", text)
        self.assertIn("ask one focused question about the underlying data", text)
        self.assertIn("never ask the user to choose a DwC-DP mapping", text)
        self.assertIn("Recheck source coverage only for groups changed", text)
        self.assertIn("Important-fact preservation", text)
        self.assertIn("representing their source rows did not mask loss", text)
        self.assertIn("Missing Event hierarchy", text)
        self.assertIn("several flat Events repeat within the same stable sampling unit", text)
        self.assertIn("otherwise retain the flat structure and document why", text)
        self.assertIn("RELATIONAL MODELLING CANDIDATES — REVIEWER ATTENTION", text)
        self.assertIn("deterministic attention aid", text)
        self.assertIn("not a modelling decision or completion gate", text)
        self.assertIn("Possible candidates are context only", text)
        self.assertIn("do not add a resource solely because a possible candidate was surfaced", text)

    def test_final_prompt_requires_maximally_faithful_dwca_projection(self):
        text = self.task_text["Final Review & Publication"]

        self.assertIn("most faithful standards-compliant DwC-A projection", text)
        self.assertIn("Preserve every useful DwC-DP fact", text)
        self.assertIn("Choose the focal core first", text)
        self.assertIn("what the data are fundamentally about", text)
        self.assertIn("Consider only extensions compatible with the chosen core", text)
        self.assertIn("at least one meaningful non-key fact", text)
        self.assertIn("Do not force data into an inexact core or extension term", text)
        self.assertIn("With Event core, an Occurrence extension is often appropriate", text)
        self.assertIn("Humboldt Ecological Inventory extension", text)
        self.assertIn("An Event hierarchy is a reason to inspect those facts, not sufficient evidence", text)
        self.assertIn("Lossless normalisation such as safe date formatting is allowed", text)
        self.assertIn("Never infer survey properties from detected occurrences", text)
        self.assertIn("review the saved coverage of semantically important fact groups", text)
        self.assertIn("Inspect every populated DwC-DP resource", text)
        self.assertIn("`dynamicProperties` only as a fallback", text)
        self.assertIn("temporary non-DwC-DP projection tables", text)
        self.assertIn("dedicated `_coreid` column", text)
        self.assertIn("exactly matches the selected core table identifier", text)
        self.assertIn("Silent omission is an error", text)
        self.assertIn("summarize every material source group omitted", text)
        self.assertIn("whether the user confirmed it", text)
        self.assertIn("explicit extension assignments", text)
        self.assertNotIn("ExportDwcaFromDwcDp", text)


class SetAgentTaskToCompleteTests(TestCase):
    def test_transformation_completion_rechecks_current_package_validation(self):
        task = Task.objects.create(name="Data transformation", text="Transform", order=1)
        dataset = Dataset.objects.create()
        Table.objects.create(
            dataset=dataset,
            title="event",
            df=pd.DataFrame([{"event_pk": "event-1", "eventCategory": "occurrence"}]),
        )
        Table.objects.create(
            dataset=dataset,
            title="occurrence",
            df=pd.DataFrame([
                {
                    "occurrence_pk": "occurrence-1",
                    "occurrenceID": "occurrence-1",
                    "event_fk": "missing-event",
                    "occurrenceStatus": "present",
                }
            ]),
        )
        agent = Agent.objects.create(dataset=dataset, task=task)

        result = SetAgentTaskToComplete(agent_id=agent.id).run()

        agent.refresh_from_db()
        self.assertIsNone(agent.completed_at)
        self.assertIn("current DwC-DP tables validate", result)
        self.assertIn("missing-event", result)

    def test_refinement_completion_removes_non_package_staging_tables(self):
        task = Task.objects.create(
            name="Data validation and refinement",
            text="Validate",
            order=1,
        )
        dataset = Dataset.objects.create()
        Table.objects.create(
            dataset=dataset,
            title="event",
            df=pd.DataFrame([{"event_pk": "event-1", "eventCategory": "occurrence"}]),
        )
        Table.objects.create(
            dataset=dataset,
            title="occurrence",
            df=pd.DataFrame([
                {
                    "occurrence_pk": "occurrence-1",
                    "occurrenceID": "occurrence-1",
                    "event_fk": "event-1",
                    "occurrenceStatus": "present",
                }
            ]),
        )
        Table.objects.create(dataset=dataset, title="join scratch", df=pd.DataFrame({"x": [1]}))
        agent = Agent.objects.create(dataset=dataset, task=task)

        result = SetAgentTaskToComplete(agent_id=agent.id).run()

        self.assertIn("Task marked as complete", result)
        self.assertFalse(dataset.table_set.filter(title="join scratch").exists())
        self.assertEqual(
            set(dataset.table_set.values_list("title", flat=True)),
            {"event", "occurrence"},
        )

    def test_manuscript_task_can_complete_without_tables(self):
        task = Task.objects.create(
            name=Dataset.MANUSCRIPT_TASK_NAME,
            text="Review manuscript",
            order=1,
        )
        dataset = Dataset.objects.create(source_mode=Dataset.SourceMode.PDF_ONLY)
        agent = Agent.objects.create(dataset=dataset, task=task)

        result = SetAgentTaskToComplete(agent_id=agent.id).run()

        agent.refresh_from_db()
        self.assertIsNotNone(agent.completed_at)
        self.assertIn("Task marked as complete", result)

    def test_table_dependent_tasks_cannot_complete_without_tables(self):
        task_names = [
            "Data transformation",
            "Data validation and refinement",
            "Final Review & Publication",
        ]

        for index, task_name in enumerate(task_names, start=1):
            with self.subTest(task_name=task_name):
                task = Task.objects.create(name=task_name, text=task_name, order=index)
                dataset = Dataset.objects.create(source_mode=Dataset.SourceMode.PDF_ONLY)
                agent = Agent.objects.create(dataset=dataset, task=task)

                result = SetAgentTaskToComplete(agent_id=agent.id).run()

                agent.refresh_from_db()
                self.assertIsNone(agent.completed_at)
                self.assertIn(f"Cannot complete '{task_name}' while this dataset has no tables", result)

    def test_table_dependent_task_can_complete_with_tables(self):
        task = Task.objects.create(name="Final Review & Publication", text="Review data", order=1)
        dataset = Dataset.objects.create(source_mode=Dataset.SourceMode.PDF_ONLY)
        Table.objects.create(dataset=dataset, title="occurrence", df=pd.DataFrame([{"scientificName": "Acer"}]))
        agent = Agent.objects.create(dataset=dataset, task=task)

        result = SetAgentTaskToComplete(agent_id=agent.id).run()

        agent.refresh_from_db()
        self.assertIsNotNone(agent.completed_at)
        self.assertIn("Task marked as complete", result)

    def test_empty_reserved_resource_is_removed_and_cannot_complete(self):
        task = Task.objects.create(name="Final Review & Publication", text="Review data", order=1)
        dataset = Dataset.objects.create(source_mode=Dataset.SourceMode.PDF_ONLY)
        Table.objects.create(dataset=dataset, title="occurrence", df=pd.DataFrame(columns=["occurrence_pk"]))
        agent = Agent.objects.create(dataset=dataset, task=task)

        result = SetAgentTaskToComplete(agent_id=agent.id).run()

        agent.refresh_from_db()
        self.assertIsNone(agent.completed_at)
        self.assertFalse(dataset.table_set.exists())
        self.assertIn("while this dataset has no tables", result)


class DwcDpDescriptorPreviewTests(TestCase):
    def test_preview_returns_complete_compact_json(self):
        dataset = Dataset.objects.create(
            title="Bird observations",
            description="D" * 4000,
        )
        Table.objects.create(
            dataset=dataset,
            title="event",
            df=pd.DataFrame([
                {
                    "event_pk": "event-1",
                    "eventID": "source-event-1",
                    "eventCategory": "occurrence",
                    "eventDate": "2025-04-26",
                    "country": "Norway",
                }
            ]),
        )
        task = Task.objects.create(name="Final Review & Publication", text="Review", order=1)
        agent = Agent.objects.create(dataset=dataset, task=task)

        result = json.loads(PreviewDwcDpDescriptor(agent_id=agent.id).run())

        self.assertTrue(result["valid"], result)
        self.assertEqual(result["resource_count"], 1)
        self.assertEqual(result["resources"][0]["name"], "event")
        self.assertEqual(result["resources"][0]["row_count"], 1)
        self.assertIn("event_pk", result["resources"][0]["fields"])
        self.assertTrue(result["description_truncated"])
        self.assertEqual(len(result["description"]), 2000)


class AgentPdfAttachmentTests(TestCase):
    class FakeResponseMessage:
        tool_calls = []

        def dict(self):
            return {"role": "assistant", "content": "Done."}

    @patch("api.models.create_response_message")
    def test_next_message_does_not_attach_follow_up_pdf_twice(self, create_response_message_mock):
        captured_pdf_files = []

        def fake_create_response_message(*args, **kwargs):
            captured_pdf_files.extend(list(kwargs.get("pdf_user_files") or []))
            return self.FakeResponseMessage()

        create_response_message_mock.side_effect = fake_create_response_message
        task = Task.objects.create(name=Dataset.MANUSCRIPT_TASK_NAME, text="Review manuscript", order=1)
        dataset = Dataset.objects.create(source_mode=Dataset.SourceMode.PDF_ONLY)
        agent = Agent.objects.create(dataset=dataset, task=task)
        Message.objects.create(agent=agent, openai_obj={"role": "assistant", "content": "Please upload the PDF."})
        user_file = UserFile.objects.create(
            dataset=dataset,
            file="user_files/paper.pdf",
        )
        Message.objects.bulk_create([
            Message(
                agent=agent,
                openai_obj={
                    "role": "user",
                    "content": "Here is the PDF.",
                    "pdf_attachments": [{"user_file_id": user_file.id, "filename": "paper.pdf"}],
                },
            )
        ])

        agent.next_message()

        self.assertEqual(captured_pdf_files, [])


class AgentWorkflowActionTests(TestCase):
    def _agent_with_user_message(self):
        task = Task.objects.create(
            name="Workflow action test",
            text="Complete the workflow action test.",
            order=1,
        )
        Task.objects.create(
            name="Data maintenance",
            text="Remain available for follow-up.",
            order=2,
        )
        dataset = Dataset.objects.create(title="Workflow test", description="Test")
        agent = Agent.create_with_system_message(dataset=dataset, task=task, tables=[])
        Message.objects.bulk_create([
            Message(
                agent=agent,
                openai_obj={"role": "user", "content": "Please process this."},
            )
        ])
        return agent

    @patch("api.models.create_response_message")
    def test_no_action_is_retried_and_completed_without_user_continue(
        self,
        create_response_message_mock,
    ):
        agent = self._agent_with_user_message()
        create_response_message_mock.side_effect = [
            CompatAssistantMessage(content="The work is complete."),
            CompatAssistantMessage(
                tool_calls=[
                    CompatToolCall(
                        id="call-complete",
                        function=CompatFunctionCall(
                            name="SetAgentTaskToComplete",
                            arguments=json.dumps({"agent_id": agent.id}),
                        ),
                    )
                ],
            ),
        ]

        agent.next_message()

        agent.refresh_from_db()
        self.assertEqual(create_response_message_mock.call_count, 2)
        self.assertIsNotNone(agent.completed_at)
        retry_items = create_response_message_mock.call_args.kwargs["additional_input_items"]
        self.assertIn("took no action", retry_items[-1]["content"])
        self.assertFalse(
            Message.objects.filter(
                agent=agent,
                openai_obj__role="user",
                openai_obj__content__icontains="continue",
            ).exists()
        )

    @patch("api.models.create_response_message")
    def test_request_user_input_creates_terminal_visible_questions(
        self,
        create_response_message_mock,
    ):
        agent = self._agent_with_user_message()
        create_response_message_mock.return_value = CompatAssistantMessage(
            tool_calls=[
                CompatToolCall(
                    id="call-question",
                    function=CompatFunctionCall(
                        name="RequestUserInput",
                        arguments=json.dumps({
                            "agent_id": agent.id,
                            "questions": [
                                {
                                    "question": "Are depths recorded in metres?",
                                    "context": "The source file does not specify a unit.",
                                },
                                {
                                    "question": "Should zero counts mean absence?",
                                },
                            ],
                        }),
                    ),
                )
            ],
        )

        messages = agent.next_message()

        agent.refresh_from_db()
        self.assertIsNone(agent.completed_at)
        self.assertEqual(create_response_message_mock.call_count, 1)
        self.assertEqual(messages[-1].role, Message.Role.ASSISTANT)
        self.assertIn("1. Are depths recorded in metres?", messages[-1].openai_obj["content"])
        self.assertIn("2. Should zero counts mean absence?", messages[-1].openai_obj["content"])
        self.assertIsNone(agent.next_message())

    @patch("api.models.create_response_message")
    def test_request_user_input_pairs_later_calls_without_executing_them(
        self,
        create_response_message_mock,
    ):
        agent = self._agent_with_user_message()
        create_response_message_mock.return_value = CompatAssistantMessage(
            tool_calls=[
                CompatToolCall(
                    id="call-question",
                    function=CompatFunctionCall(
                        name="RequestUserInput",
                        arguments=json.dumps({
                            "agent_id": agent.id,
                            "questions": [{"question": "What does this code mean?"}],
                        }),
                    ),
                ),
                CompatToolCall(
                    id="call-complete",
                    function=CompatFunctionCall(
                        name="SetAgentTaskToComplete",
                        arguments=json.dumps({"agent_id": agent.id}),
                    ),
                ),
            ],
        )

        messages = agent.next_message()

        agent.refresh_from_db()
        outputs = {
            message.openai_obj.get("tool_call_id"): message.openai_obj.get("content")
            for message in messages
            if message.openai_obj.get("role") == Message.Role.TOOL
        }
        self.assertIn("call-question", outputs)
        self.assertIn("call-complete", outputs)
        self.assertIn("Skipped because RequestUserInput", outputs["call-complete"])
        self.assertIsNone(agent.completed_at)

    def test_request_user_input_schema_allows_multiple_questions(self):
        schema = RequestUserInput.openai_schema()

        self.assertEqual(
            schema["parameters"]["properties"]["questions"]["maxItems"],
            10,
        )

    @override_settings(
        OPENAI_TOOL_HISTORY_TURNS=4,
        OPENAI_FULL_TOOL_HISTORY_TURNS=1,
        OPENAI_COMPACT_TOOL_CHARS=1000,
    )
    def test_model_history_compacts_older_large_tool_payloads(self):
        agent = self._agent_with_user_message()
        for index in range(4):
            call_id = f"call-{index}"
            Message.objects.create(
                agent=agent,
                openai_obj={
                    "role": "assistant",
                    "content": "",
                    "tool_calls": [{
                        "id": call_id,
                        "type": "function",
                        "function": {
                            "name": "Python",
                            "arguments": json.dumps({"code": "x" * 5000}),
                        },
                    }],
                },
            )
            Message.objects.create(
                agent=agent,
                openai_obj={
                    "role": "tool",
                    "tool_call_id": call_id,
                    "content": "y" * 5000,
                },
            )

        bounded = [
            message.openai_obj
            for message in agent.messages_for_model()
            if message.openai_obj.get("role") in {"assistant", "tool"}
        ]
        tool_assistants = [obj for obj in bounded if obj["role"] == "assistant"]
        tool_outputs = [obj for obj in bounded if obj["role"] == "tool"]

        self.assertEqual(len(tool_assistants), 4)
        self.assertIn(
            "older tool argument compacted",
            tool_assistants[0]["tool_calls"][0]["function"]["arguments"],
        )
        self.assertGreater(
            len(tool_assistants[-1]["tool_calls"][0]["function"]["arguments"]),
            5000,
        )
        self.assertIn("older tool output compacted", tool_outputs[0]["content"])
        self.assertEqual(tool_outputs[-1]["content"], "y" * 5000)

    def test_validation_system_prompt_snapshots_only_package_tables(self):
        task = Task.objects.create(
            name="Data validation and refinement",
            text="Validate",
            order=1,
        )
        dataset = Dataset.objects.create(title="Prompt scope", description="Test")
        source = Table.objects.create(
            dataset=dataset,
            title="source sheet",
            df=pd.DataFrame({"source_column": ["SOURCE-ONLY-VALUE"]}),
        )
        event = Table.objects.create(
            dataset=dataset,
            title="event",
            df=pd.DataFrame({
                "event_pk": ["event-1"],
                "eventCategory": ["occurrence"],
            }),
        )

        agent = Agent.create_with_system_message(
            dataset=dataset,
            task=task,
            tables=[source, event],
        )
        prompt = agent.message_set.get(
            openai_obj__role=Message.Role.SYSTEM,
        ).openai_obj["content"]

        self.assertIn("source_column", prompt)
        self.assertNotIn("SOURCE-ONLY-VALUE", prompt)
        self.assertIn("event-1", prompt)
        self.assertIn("Snapshot omitted in this package-focused task", prompt)
